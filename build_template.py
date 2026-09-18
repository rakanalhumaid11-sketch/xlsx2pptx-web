"""يبني ملف القالب الثابت (template.pptx) من تقرير سابق جاهز.

يُشغَّل مرة واحدة فقط عند الحاجة لتحديث القالب. الناتج: قالب من 4 شرائح
(الغلاف، الملخص + الجدول، شريحة ملاحظة نموذجية فارغة، شكرًا لكم) وكل شكل
فيه يحمل اسمًا واضحًا (PH_*) ليجده المولّد لاحقًا بالاسم لا بالموقع.
"""
import io
import sys

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, "/home/claude/xlsx2pptx_web")
import engine

SRC = "/mnt/user-data/uploads/تقرير_الملاحظات (1).pptx"
DST = "/home/claude/xlsx2pptx_web/template.pptx"

# صفوف الجدول في شريحة الملخص: عمودان (يمين ثم يسار لأن الاتجاه عربي)
RIGHT_COUNT_L, RIGHT_DESC_L = 6246000, 7146000
LEFT_COUNT_L, LEFT_DESC_L = 400000, 1300000
ROW_TOPS_RIGHT = [2560000 + 360000 * i for i in range(11)]
ROW_TOPS_LEFT = [2560000 + 360000 * i for i in range(10)]


def by_id(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"shape id {shape_id} not found")


def at(slide, left, top):
    for sh in slide.shapes:
        if sh.left == left and sh.top == top:
            return sh
    raise KeyError(f"shape at {left},{top} not found")


def find_named(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def copy_text_body(src, dst):
    """ينسخ بنية النص (الخط والحجم واللون والمحاذاة) من شكل إلى آخر."""
    from copy import deepcopy
    from pptx.oxml.ns import qn
    src_tx = src._element.find(qn("p:txBody"))
    dst_tx = dst._element.find(qn("p:txBody"))
    if src_tx is None or dst_tx is None:
        return
    dst._element.replace(dst_tx, deepcopy(src_tx))


def placeholder_png(w=600, h=800, shade=0xE8):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (shade, shade, shade)).save(buf, format="PNG")
    return buf.getvalue()


def main():
    prs = Presentation(SRC)
    n = len(prs.slides)
    keep = [0, 1, 2, n - 1]
    drop = [i for i in range(n) if i not in keep]
    engine.rebuild_slide_order(prs, keep, drop)

    cover, summary, note, thanks = (prs.slides[i] for i in range(4))

    # ---------- الغلاف ----------
    t = by_id(cover, 2)
    t.name = "PH_COVER_TITLE"          # سطر 1 ثابت، سطر 2 = اسم المغذي
    d = by_id(cover, 5)
    d.name = "PH_COVER_DATE"

    # ---------- الملخص ----------
    by_id(summary, 11).name = "PH_SUM_HEADING"
    by_id(summary, 12).name = "PH_SUM_SUBTITLE"
    by_id(summary, 13).name = "PH_KPI_BOX_TOTAL"
    by_id(summary, 14).name = "PH_KPI_TOTAL"
    by_id(summary, 15).name = "PH_KPI_LBL_TOTAL"
    by_id(summary, 16).name = "PH_KPI_BOX_TYPES"
    by_id(summary, 17).name = "PH_KPI_TYPES"
    by_id(summary, 18).name = "PH_KPI_LBL_TYPES"
    by_id(summary, 19).name = "PH_KPI_BOX_DONE"
    by_id(summary, 20).name = "PH_KPI_DONE"
    by_id(summary, 21).name = "PH_KPI_LBL_DONE"
    by_id(summary, 22).name = "PH_TH_R_COUNT"
    by_id(summary, 23).name = "PH_TH_R_DESC"
    by_id(summary, 46).name = "PH_TH_L_COUNT"
    by_id(summary, 47).name = "PH_TH_L_DESC"

    for i, top in enumerate(ROW_TOPS_RIGHT):
        c = at(summary, RIGHT_COUNT_L, top)
        c.name = f"PH_ROW_{i:02d}_COUNT"
        engine.set_shape_text_preserve_style(c, "")
        d2 = at(summary, RIGHT_DESC_L, top)
        d2.name = f"PH_ROW_{i:02d}_DESC"
        engine.set_shape_text_preserve_style(d2, "")
    for j, top in enumerate(ROW_TOPS_LEFT):
        i = 11 + j
        c = at(summary, LEFT_COUNT_L, top)
        c.name = f"PH_ROW_{i:02d}_COUNT"
        engine.set_shape_text_preserve_style(c, "")
        d2 = at(summary, LEFT_DESC_L, top)
        d2.name = f"PH_ROW_{i:02d}_DESC"
        engine.set_shape_text_preserve_style(d2, "")

    # صفوف الجدول التي كانت فارغة في التقرير الأصلي لا تحمل تنسيق خط (حجم/لون)،
    # فلو كُتب فيها نص لظهر بخط أكبر ومختلف عن بقية الصفوف. ننسخ لها نفس بنية
    # النص من صف مكتمل التنسيق (مع إبقاء موضعها ولون خلفيتها كما هما).
    for i in (18, 19, 20):
        for suffix, src_idx in ((("COUNT"), 11), ("DESC", 11)):
            src = find_named(summary, f"PH_ROW_{src_idx:02d}_{suffix}")
            dst = find_named(summary, f"PH_ROW_{i:02d}_{suffix}")
            if src is not None and dst is not None:
                copy_text_body(src, dst)

    for sid, txt in ((14, ""), (17, ""), (20, "")):
        engine.set_shape_text_preserve_style(by_id(summary, sid), txt)

    # ---------- شريحة الملاحظة ----------
    by_id(note, 2).name = "PH_NOTE_HEADER"
    by_id(note, 20).name = "PH_NOTICE"
    by_id(note, 21).name = "PH_NOTICE_LABEL"
    by_id(note, 7).name = "PH_NOTE_TITLE"
    by_id(note, 9).name = "PH_OFFICE"
    by_id(note, 13).name = "PH_COORDS"
    by_id(note, 15).name = "PH_DESC"
    by_id(note, 11).name = "PH_LBL_BEFORE"
    by_id(note, 12).name = "PH_LBL_AFTER"
    by_id(note, 3).name = "PH_FRAME_BEFORE"
    by_id(note, 5).name = "PH_FRAME_AFTER"

    for sid in (20, 7, 9, 13, 15):
        engine.set_shape_text_preserve_style(by_id(note, sid), "")

    # صورة "بعد" (اليسار) تُحذف نهائيًا: الخانة تبقى فارغة دائمًا
    engine.remove_shape(by_id(note, 38))
    # صورة "قبل" (اليمين) تبقى كعنصر نائب رمادي صغير يُستبدل بصورة كل ملاحظة
    before = by_id(note, 39)
    newpic = engine.replace_picture_shape(note, before, placeholder_png())
    newpic.name = "PH_PHOTO_BEFORE"

    by_id(thanks, 8).name = "PH_THANKS"

    prs.save(DST)
    print("saved", DST)

    chk = Presentation(DST)
    print("slides:", len(chk.slides))
    for i, s in enumerate(chk.slides):
        names = [sh.name for sh in s.shapes if sh.name.startswith("PH_")]
        print(f"  slide {i}: {len(s.shapes)} shapes, {len(names)} named")
    import os
    print("size:", round(os.path.getsize(DST) / 1024, 1), "KB")


if __name__ == "__main__":
    main()
