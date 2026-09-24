"""مولّد تقرير صيانة المغذي من ملف إكسل واحد، باستخدام قالب ثابت مخزّن.

لا يحتاج المستخدم لاختيار قالب ولا لتحديد أعمدة ولا لعدد الشرائح: القالب
(template.pptx) مخزّن بجانب هذا الملف، وكل شكل فيه يحمل اسمًا واضحًا يبدأ
بـ PH_ فيجده المولّد بالاسم (وليس بالموقع، حتى لا يتأثر بأي تحريك بسيط).

الشرائح الثابتة: الغلاف، الملخص + الجدول (يتوسّع لعدة شرائح عند الحاجة)،
ثم شريحة لكل ملاحظة، ثم شريحة الشكر.
"""
from __future__ import annotations

import datetime
import gc
import os
import re
import shutil
import zipfile
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pptx import Presentation

import engine

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "template.pptx")

IDX_COVER, IDX_SUMMARY, IDX_NOTE, IDX_THANKS = 0, 1, 2, 3

# سعة جدول الملخص في الشريحة الواحدة: 11 صفًا في العمود الأيمن + 10 في الأيسر
ROWS_PER_SUMMARY = 21

# عدد الملاحظات التي تُجهَّز صورها دفعة واحدة
CHUNK = 16

MONTHS_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
             "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]

# الأعمدة المطلوبة من الإكسل: المفتاح الداخلي -> الأسماء المحتملة في الترويسة
COLUMN_ALIASES: Dict[str, List[str]] = {
    "note_id": ["رقم الملاحظة", "رقم الملاحظه", "الملاحظة رقم"],
    "station": ["المحطة", "المحطه", "رقم المحطة"],
    "feeder": ["المغذي", "المغذى", "اسم المغذي"],
    "note": ["الملاحظة", "الملاحظه", "وصف الملاحظة"],
    "office": ["المكتب"],
    "lat": ["خط العرض", "دائرة العرض"],
    "lon": ["خط الطول"],
    # النظام غيّر تسمية أعمدة الصور من «صورة 1» إلى «صورة قبل 1»، ولولا
    # إضافتها هنا لخرج التقرير بلا صور ولخلت دبابيس الخريطة من صورة الملاحظة
    "photo1": ["صورة 1", "صورة1", "الصورة 1", "صوره 1",
               "صورة قبل 1", "صوره قبل 1", "صورة قبل1"],
    "photo2": ["صورة 2", "صورة2", "الصورة 2", "صوره 2",
               "صورة قبل 2", "صوره قبل 2", "صورة قبل2"],
    # صار النظام يصدّر صور ما بعد المعالجة أيضًا، فتُملأ خانة «بعد» من الملف
    "photo_after": ["صورة بعد 1", "صوره بعد 1", "صورة بعد1"],
    "status": ["حالة الملاحظة", "حالة الملاحظه", "الحالة"],
    "inspect_date": ["تاريخ الفحص", "التاريخ"],
    "contractor": ["المقاول", "اسم المقاول", "الشركة", "الشركه", "المنفذ",
                   "تمت المعالجة بواسطة", "تمت المعالجه بواسطة"],
}


# ---------------------------------------------------------------- قراءة الإكسل

def _clean(v: Any) -> str:
    """يحوّل قيمة خلية إلى نص نظيف، ويتجاهل أخطاء إكسل مثل ‎#N/A‎ و ‎#VALUE!‎."""
    if v is None:
        return ""
    s = str(v).strip()
    if not s or s.startswith("#") and s.upper() in ("#N/A", "#VALUE!", "#REF!", "#NAME?", "#DIV/0!"):
        return ""
    return s


def _match_column(header: str, aliases: List[str]) -> bool:
    h = engine.normalize_ar(header)
    for a in aliases:
        na = engine.normalize_ar(a)
        if h == na or (na and h.startswith(na)):
            return True
    return False


def _map_headers(headers: List[str]) -> Dict[str, int]:
    """يربط كل مفتاح داخلي برقم عموده في الإكسل اعتمادًا على اسم الترويسة."""
    mapping: Dict[str, int] = {}
    for key, aliases in COLUMN_ALIASES.items():
        for idx, h in enumerate(headers):
            if not h:
                continue
            if _match_column(h, aliases):
                # "الملاحظة" قد تطابق "رقم الملاحظة" أيضًا، لذا نفضّل المطابقة التامة
                if key in mapping:
                    if engine.normalize_ar(h) in [engine.normalize_ar(a) for a in aliases]:
                        mapping[key] = idx
                else:
                    mapping[key] = idx
    return mapping


def read_excel(path: str) -> Tuple[List[Dict[str, str]], Dict[str, int], List[str]]:
    """يقرأ أفضل ورقة في الملف ويرجع (السجلات، خريطة الأعمدة، الترويسة)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    best: Optional[Tuple[int, str, int, List[str]]] = None
    for name in wb.sheetnames:
        ws = wb[name]
        try:
            hdr_idx = engine._detect_header_row(ws)
        except Exception:
            hdr_idx = 0
        headers: List[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == hdr_idx:
                headers = [_clean(c) for c in row]
                break
        mapping = _map_headers(headers)
        score = len(mapping)
        if best is None or score > best[0]:
            best = (score, name, hdr_idx, headers)
    if best is None:
        raise ValueError("الملف لا يحتوي على أوراق عمل.")

    _, sheet_name, hdr_idx, headers = best
    mapping = _map_headers(headers)
    if "note" not in mapping:
        raise ValueError("لم يُعثر على عمود «الملاحظة» في الإكسل.")

    ws = wb[sheet_name]
    records: List[Dict[str, str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= hdr_idx:
            continue
        rec = {}
        for key, col in mapping.items():
            rec[key] = _clean(row[col]) if col < len(row) else ""
        if not any(rec.values()):
            continue
        if not rec.get("note") and not rec.get("note_id"):
            continue
        records.append(rec)
    wb.close()
    return records, mapping, headers


# -------------------------------------------------- الصور المخزّنة داخل الإكسل

_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_WEBIMG_NS = "{http://schemas.microsoft.com/office/spreadsheetml/2020/richdatawebimage}"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

WEBIMG_XML = "xl/richData/rdRichValueWebImage.xml"
WEBIMG_RELS = "xl/richData/_rels/rdRichValueWebImage.xml.rels"


def _embedded_image_map(zf: zipfile.ZipFile) -> Dict[str, str]:
    """يبني خريطة {رابط الصورة -> مسار الصورة داخل ملف الإكسل}.

    ملفات إكسل المصدَّرة من النظام تخزّن نسخة من كل صورة داخلها («صور في
    الخلايا»)، مربوطة برابطها الأصلي. استخدام هذه النسخة يوفّر تنزيل مئات
    الصور من الإنترنت في كل مرة، فيصبح التوليد أسرع بكثير ولا يتأثر بانقطاع
    الشبكة أو بطئها."""
    names = set(zf.namelist())
    if WEBIMG_XML not in names or WEBIMG_RELS not in names:
        return {}
    try:
        rels = {}
        for rel in ET.fromstring(zf.read(WEBIMG_RELS)):
            rid, target = rel.get("Id"), rel.get("Target", "")
            if rid:
                rels[rid] = target
        out: Dict[str, str] = {}
        for srd in ET.fromstring(zf.read(WEBIMG_XML)):
            addr = srd.find(f"{_WEBIMG_NS}address")
            blip = srd.find(f"{_WEBIMG_NS}blip")
            if addr is None or blip is None:
                continue
            url = rels.get(addr.get(_R_ID, ""), "")
            media = rels.get(blip.get(_R_ID, ""), "")
            if not url or not media:
                continue
            path = media.replace("../", "xl/").lstrip("/")
            if path in names:
                out[url.strip()] = path
        return out
    except (ET.ParseError, KeyError, OSError):
        return {}


class ImageSource:
    """يجلب صورة الملاحظة: من داخل ملف الإكسل إن وُجدت، وإلا ينزّلها من رابطها.

    تُقرأ كل صورة عند الحاجة إليها فقط ثم تُهمل، فلا تتراكم الصور في الذاكرة
    (كانت هذه سبب توقف التوليد سابقًا على الاستضافة المجانية)."""

    def __init__(self, xlsx_path: str):
        self.zf: Optional[zipfile.ZipFile] = None
        self.map: Dict[str, str] = {}
        self.from_file = 0
        self.from_web = 0
        self.max_workers = int(os.environ.get("IMAGE_WORKERS", "8"))
        try:
            self.zf = zipfile.ZipFile(xlsx_path)
            self.map = _embedded_image_map(self.zf)
        except (zipfile.BadZipFile, OSError):
            self.zf = None

    def get(self, url: str) -> Optional[bytes]:
        url = (url or "").strip()
        if not url:
            return None
        path = self.map.get(url)
        if path and self.zf is not None:
            try:
                data = self.zf.read(path)
                if data:
                    self.from_file += 1
                    return data
            except (KeyError, OSError):
                pass
        if url.startswith("http"):
            data = engine.download_image(url)
            if data:
                self.from_web += 1
            return data
        return None

    def get_many(self, urls: List[str]) -> Dict[str, Optional[bytes]]:
        """يجلب مجموعة صور دفعة واحدة. الصور المخزّنة داخل الإكسل تُقرأ فورًا،
        وما يحتاج تنزيلًا من الإنترنت يُنزَّل بالتوازي: تنزيل ستمئة صورة واحدة
        تلو الأخرى قد يستغرق نصف ساعة، وبالتوازي دقائق معدودة."""
        out: Dict[str, Optional[bytes]] = {}
        need_web = []
        for url in urls:
            url = (url or "").strip()
            if not url or url in out:
                continue
            path = self.map.get(url)
            if path and self.zf is not None:
                try:
                    data = self.zf.read(path)
                except (KeyError, OSError):
                    data = None
                if data:
                    self.from_file += 1
                    out[url] = data
                    continue
            if url.startswith("http"):
                need_web.append(url)
            else:
                out[url] = None

        if need_web:
            from concurrent.futures import ThreadPoolExecutor
            workers = min(self.max_workers, len(need_web))
            with ThreadPoolExecutor(max_workers=workers) as pool:
                for url, data in zip(need_web, pool.map(engine.download_image, need_web)):
                    out[url] = data
                    if data:
                        self.from_web += 1
        return out

    def close(self):
        if self.zf is not None:
            try:
                self.zf.close()
            except OSError:
                pass
            self.zf = None


# ------------------------------------------------------------- اشتقاق البيانات

def feeder_code(records: List[Dict[str, str]]) -> str:
    """يبني رمز المغذي مثل «M13-8903»: الرمز اللاتيني من عمود المغذي + رقم المحطة."""
    prefix, station = "", ""
    for rec in records:
        if not station and rec.get("station"):
            station = rec["station"].split(".")[0]
        raw = rec.get("feeder", "")
        if not prefix and raw:
            m = re.search(r"\(\s*([A-Za-z]{1,4}\s*\d{1,4})\s*[-–]", raw)
            if not m:
                m = re.search(r"([A-Za-z]{1,4}\s*\d{1,4})", raw)
            if m:
                prefix = m.group(1).replace(" ", "").upper()
        if prefix and station:
            break
    if prefix and station:
        return f"{prefix}-{station}"
    return station or prefix or ""


def today_ar() -> str:
    d = datetime.date.today()
    return f"{d.day} {MONTHS_AR[d.month - 1]}، {d.year}"


def note_counts(records: List[Dict[str, str]]) -> List[Tuple[str, int]]:
    c = Counter(r["note"] for r in records if r.get("note"))
    return sorted(c.items(), key=lambda kv: (-kv[1], kv[0]))


# ----------------------------------------------------------- أدوات على الشرائح

def find(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def set_text(slide, name: str, text: str):
    sh = find(slide, name)
    if sh is not None and sh.has_text_frame:
        engine.set_shape_text_preserve_style(sh, text)
    return sh


def set_paragraph(shape, index: int, text: str):
    """يغيّر نص فقرة واحدة داخل الشكل مع الحفاظ على تنسيقها (للغلاف ذي السطرين)."""
    tf = shape.text_frame
    if index >= len(tf.paragraphs):
        return
    p = tf.paragraphs[index]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.text = text


def drop(slide, name: str):
    sh = find(slide, name)
    if sh is not None:
        engine.remove_shape(sh)


# -------------------------------------------------------------- تعبئة الشرائح

def fill_cover(slide, feeder: str, date_text: str):
    title = find(slide, "PH_COVER_TITLE")
    if title is not None:
        # عنوان الغلاف سطران: «صيانة مغذي» ثم رمز المغذي. نستهدف آخر فقرة فيها
        # نص فعلي (الصندوق يحوي فقرات فارغة للمباعدة يجب ألا نكتب فيها).
        paras = title.text_frame.paragraphs
        target = None
        for i, p in enumerate(paras):
            if p.text.strip():
                target = i
        if target is not None:
            set_paragraph(title, target, feeder)
    set_text(slide, "PH_COVER_DATE", date_text)


def fill_summary(slide, feeder: str, notice: str, totals: Optional[Dict[str, int]],
                 rows: List[Tuple[str, int]], continued: bool, extra: str = ""):
    """يملأ شريحة ملخص واحدة. `rows` جزء الجدول الخاص بهذه الشريحة فقط.

    في شرائح التكملة نحذف مربعات الأرقام الثلاثة حتى لا تتكرر نفس القيم."""
    subtitle = f"صيانة مغذي: {feeder}" if feeder else "صيانة مغذي"
    if notice:
        subtitle += f"  |  رقم اشعار فحص : {notice}"
    if extra:
        subtitle += f"  |  {extra}"
    set_text(slide, "PH_SUM_SUBTITLE", subtitle)
    if continued:
        set_text(slide, "PH_SUM_HEADING", "ملخص الملاحظات (تابع)")
        for n in ("PH_KPI_BOX_TOTAL", "PH_KPI_TOTAL", "PH_KPI_LBL_TOTAL",
                  "PH_KPI_BOX_TYPES", "PH_KPI_TYPES", "PH_KPI_LBL_TYPES",
                  "PH_KPI_BOX_DONE", "PH_KPI_DONE", "PH_KPI_LBL_DONE"):
            drop(slide, n)
    elif totals:
        set_text(slide, "PH_KPI_TOTAL", str(totals["total"]))
        set_text(slide, "PH_KPI_TYPES", str(totals["types"]))
        set_text(slide, "PH_KPI_DONE", str(totals["done"]))

    for i in range(ROWS_PER_SUMMARY):
        cname, dname = f"PH_ROW_{i:02d}_COUNT", f"PH_ROW_{i:02d}_DESC"
        if i < len(rows):
            desc, cnt = rows[i]
            set_text(slide, cname, str(cnt))
            set_text(slide, dname, desc)
        else:
            # الصفوف غير المستخدمة تُحذف حتى لا تظهر خانات فارغة في الجدول
            drop(slide, cname)
            drop(slide, dname)

    # إذا لم يُستخدم العمود الأيسر إطلاقًا نحذف ترويسته أيضًا
    if len(rows) <= 11:
        drop(slide, "PH_TH_L_COUNT")
        drop(slide, "PH_TH_L_DESC")


_UNSET = object()


def fill_note(slide, rec: Dict[str, str], feeder: str, notice: str,
              images: Optional["ImageSource"] = None, img=_UNSET,
              img_after=_UNSET) -> Tuple[bool, bool]:
    """يملأ شريحة ملاحظة واحدة بصورتَي «قبل» و«بعد». يرجع (وُضعت قبل، وُضعت بعد)."""
    header = "صيانة المغذي"
    if feeder:
        header += f" {feeder}"
    set_text(slide, "PH_NOTE_HEADER", header)
    set_text(slide, "PH_NOTICE", notice)

    nid = rec.get("note_id", "")
    set_text(slide, "PH_NOTE_TITLE", f"معالجة الملاحظة رقم {nid}" if nid else "معالجة الملاحظة")
    set_text(slide, "PH_OFFICE", rec.get("office", ""))
    set_text(slide, "PH_DESC", rec.get("note", ""))

    lat, lon = rec.get("lat", ""), rec.get("lon", "")
    coords_shape = find(slide, "PH_COORDS")
    if coords_shape is not None:
        text = f"{lat}, {lon}" if lat and lon else (lat or lon)
        engine.set_shape_text_preserve_style(coords_shape, text)
        if lat and lon:
            url = engine.maps_url(lat, lon)
            if url:
                engine.set_hyperlink(coords_shape, url)

    def fetch(url, cached):
        if cached is not _UNSET:
            return cached
        if images is not None:
            return images.get(url)
        return engine.download_image(url) if url.startswith("http") else None

    # خانة «بعد»: إطار بلا شكل صورة، فنُدرج الصورة داخله عند توفّرها في الملف
    after_ok = False
    frame = find(slide, "PH_FRAME_AFTER")
    after = fetch(rec.get("photo_after", ""), img_after)
    if frame is not None and after:
        engine.insert_picture_in_box(slide, frame.left, frame.top,
                                     frame.width, frame.height, after)
        after_ok = True

    box = find(slide, "PH_PHOTO_BEFORE")
    if box is None:
        return False, after_ok
    img = fetch(rec.get("photo1", ""), img)
    if img:
        engine.replace_picture_shape(slide, box, img)
        return True, after_ok
    # لا صورة متاحة: نحذف العنصر النائب فتبقى الخانة فارغة بدل تكرار صورة القالب
    engine.remove_shape(box)
    return False, after_ok


# --------------------------------------------------------------- نقطة الدخول

def _build_one(out_path: str, part_records: List[Dict[str, str]], feeder: str,
               notice: str, counts: List[Tuple[str, int]], totals: Dict[str, int],
               images: "ImageSource", part: Tuple[int, int],
               progress_cb=None, offset: int = 0, grand_total: int = 0) -> Dict[str, int]:
    """يبني ملفًا واحدًا لمجموعة ملاحظات، ويحفظه ثم يتركه للذاكرة أن تتحرر.

    الملخص والأرقام تخصّ المغذي كاملًا في كل جزء (لا جزءه فقط)، لأنها حقيقة
    واحدة عن المغذي؛ وما يختلف بين الأجزاء هو شرائح الملاحظات."""
    prs = Presentation(TEMPLATE_PATH)
    n_summary = max(1, (len(counts) + ROWS_PER_SUMMARY - 1) // ROWS_PER_SUMMARY)

    part_no, part_count = part
    label = f"الجزء {part_no} من {part_count}" if part_count > 1 else ""
    date_text = today_ar() + (f" · {label}" if label else "")
    fill_cover(prs.slides[IDX_COVER], feeder, date_text)

    # شرائح الملخص: الأولى موجودة في القالب، والباقي نسخ منها
    summary_indices = [IDX_SUMMARY]
    for _ in range(n_summary - 1):
        summary_indices.append(_dup(prs, IDX_SUMMARY))
    for k, si in enumerate(summary_indices):
        chunk = counts[k * ROWS_PER_SUMMARY:(k + 1) * ROWS_PER_SUMMARY]
        fill_summary(prs.slides[si], feeder, notice, totals, chunk,
                     continued=(k > 0), extra=label)

    note_indices: List[int] = []
    photos_ok = 0
    after_ok = 0
    total = len(part_records)
    # نعالج على دفعات: ننزّل صور الدفعة بالتوازي ثم نبني شرائحها ونتخلص منها،
    # فلا تتراكم الصور في الذاكرة مهما كان عدد الملاحظات.
    for start in range(0, total, CHUNK):
        chunk = part_records[start:start + CHUNK]
        blobs = images.get_many([r.get(k, "") for r in chunk
                                 for k in ("photo1", "photo_after")])
        for j, rec in enumerate(chunk):
            idx = _dup(prs, IDX_NOTE)
            note_indices.append(idx)
            img = blobs.get((rec.get("photo1", "") or "").strip())
            aft = blobs.get((rec.get("photo_after", "") or "").strip())
            before, after = fill_note(prs.slides[idx], rec, feeder, notice,
                                      images, img=img, img_after=aft)
            photos_ok += before
            after_ok += after
            if progress_cb:
                progress_cb(offset + start + j + 1, grand_total or total)
        blobs.clear()

    final_order = [IDX_COVER] + summary_indices + note_indices + [IDX_THANKS]
    engine.rebuild_slide_order(prs, final_order, [IDX_NOTE])
    prs.save(out_path)
    return {"slides": len(final_order), "photos_ok": photos_ok,
            "after_ok": after_ok, "summary_slides": n_summary}


def build_report(excel_path: str, out_path: str, notice: str = "",
                 progress_cb=None) -> Dict[str, Any]:
    records, mapping, headers = read_excel(excel_path)
    if not records:
        raise ValueError("لم يُعثر على أي ملاحظات في ملف الإكسل.")

    feeder = feeder_code(records)
    counts = note_counts(records)
    totals = {"total": len(records), "types": len(counts), "done": 0}

    images = ImageSource(excel_path)
    if progress_cb:
        # نُبلّغ مصدر الصور مبكرًا: القراءة من داخل الملف تستغرق ثوانٍ، أما
        # التنزيل من الإنترنت فدقائق — والمستخدم يجب أن يعرف الفرق بدل أن
        # يظن أن التوليد متوقف.
        embedded = sum(1 for r in records[:20] if r.get("photo1") in images.map)
        progress_cb(0, len(records), "file" if embedded > 10 else "web")

    # المخرج ملف بوربوينت واحد دائمًا مهما كبر عدد الملاحظات. الذروة مقيسة:
    # نحو 190 ميجا ثابتة لبنية 600 شريحة + حجم الصور بعد التصغير، أي ~285 ميجا
    # لستمئة ملاحظة — دون سقف الخادم (512) بهامش مريح.
    try:
        st = _build_one(out_path, records, feeder, notice, counts, totals,
                        images, (1, 1), progress_cb, 0, len(records))
    finally:
        images.close()
    gc.collect()

    return {
        "records": len(records),
        "types": len(counts),
        "slides": st["slides"],
        "summary_slides": st["summary_slides"],
        "photos_ok": st["photos_ok"],
        "photos_missing": len(records) - st["photos_ok"],
        "after_ok": st["after_ok"],
        "photos_from_file": images.from_file,
        "photos_from_web": images.from_web,
        "feeder": feeder,
        "columns_found": sorted(mapping.keys()),
    }


def _dup(prs: Presentation, index: int) -> int:
    engine.duplicate_slide(prs, index)
    return len(prs.slides._sldIdLst) - 1
