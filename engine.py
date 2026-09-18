# -*- coding: utf-8 -*-
"""
engine.py
==========
المحرك الأساسي لأداة "تحويل الإكسل إلى بوربوينت".
هذا الملف لا يتفاعل مع المستخدم مباشرة (ذلك عمل wizard.py) بل يوفر
كل الدوال اللازمة لـ:
  1) قراءة ملف الإكسل واقتراح ربط الأعمدة بالحقول المطلوبة.
  2) تحليل قالب البوربوينت واكتشاف الأشكال "الثابتة" مقابل "المتغيرة".
  3) استنساخ السلايد النموذجي وتعبئته بالبيانات + تنزيل وإدراج الصور.
  4) تحديث سلايد الملخص الإحصائي (Best effort).
  5) حفظ/تحميل إعدادات الربط (mapping.json) لإعادة استخدامها لاحقًا.
"""

import copy
import difflib
import io
import json
import os
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import requests
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from pptx.parts.image import Image as PptxImage
from pptx.enum.text import MSO_AUTO_SIZE


# --------------------------------------------------------------------------
# أدوات نصية عامة
# --------------------------------------------------------------------------

_TASHKEEL = re.compile(r"[ؗ-ًؚ-ْٰـ]")


def normalize_ar(text: Optional[str]) -> str:
    """توحيد شكل النص العربي لتسهيل المطابقة: إزالة التشكيل، توحيد الألف
    والهمزات والتاء المربوطة، إزالة المسافات الزائدة، تحويل الأرقام
    العربية إلى إنجليزية."""
    if text is None:
        return ""
    text = str(text)
    text = unicodedata.normalize("NFKC", text)
    text = _TASHKEEL.sub("", text)
    text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    text = text.replace("ى", "ي").replace("ة", "ه")
    text = text.replace("ؤ", "و").replace("ئ", "ي")
    # أرقام عربية -> إنجليزية
    arabic_digits = "٠١٢٣٤٥٦٧٨٩"
    for i, d in enumerate(arabic_digits):
        text = text.replace(d, str(i))
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def similarity(a: str, b: str) -> float:
    a, b = normalize_ar(a), normalize_ar(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.9
    return difflib.SequenceMatcher(None, a, b).ratio()


# --------------------------------------------------------------------------
# 1) قراءة الإكسل + اقتراح ربط الأعمدة
# --------------------------------------------------------------------------

# كل حقل مطلوب + المرادفات العربية المحتملة لاسم عموده في ملفات مختلفة
FIELD_ALIASES: Dict[str, List[str]] = {
    "seq":               ["م", "الرقم", "رقم تسلسلي", "#"],
    "note_id":           ["رقم الملاحظة", "رقم ملاحظة", "note id", "id"],
    "station":           ["المحطة", "محطة"],
    "feeder":            ["المغذي", "مغذي", "feeder"],
    "note_text":         ["الملاحظة", "نص الملاحظة", "وصف الملاحظة", "الوصف"],
    "type_code":         ["كود", "رمز التصنيف"],
    "main_category":     ["التصنيف الرئيسي", "التصنيف"],
    "priority":          ["الاولوية", "الأولوية", "priority"],
    "inspection_date":   ["تاريخ الفحص", "التاريخ", "date"],
    "isolation_point":   ["نقطة العزل", "نقطة عزل"],
    "nearest_isolation": ["نقطة العزل الاقرب", "اقرب نقطة عزل"],
    "longitude":         ["خط الطول", "longitude", "lng", "long"],
    "latitude":          ["خط العرض", "latitude", "lat"],
    "inspection_notes":  ["ملاحظات الفحص"],
    "city":              ["الادارة", "المدينة", "الادارة (المدينة)", "city"],
    "office":            ["المكتب", "office"],
    "photo_before":      ["صورة 1", "صورة قبل", "الصورة الاولى", "photo 1", "photo before"],
    "photo_after":       ["صورة 2", "صورة بعد", "الصورة الثانية", "photo 2", "photo after"],
    "note_classification": ["تصنيف الملاحظة"],
    "note_status":       ["حالة الملاحظة", "المعالجة", "حالة المعالجة"],
    "contractor":        ["المقاول", "contractor"],
    "notice_number":     ["رقم اشعار", "رقم إشعار", "رقم اشعار الصيانة", "رقم الاشعار"],
}

# الحقول الأساسية التي يفضّل إيجادها لضمان عمل الأداة بشكل جيد
CORE_FIELDS = [
    "seq", "note_text", "office", "latitude", "longitude",
    "photo_before", "photo_after",
]


@dataclass
class SheetInfo:
    name: str
    header_row_idx: int
    headers: List[str]
    n_rows: int


def _is_text_like(v) -> bool:
    """True لو القيمة نص وصفي (تسمية عمود محتملة) لا رقم صرف. صف الرؤوس
    الحقيقي يكون شبه كامل بنصوص من هذا النوع، بعكس صفوف البيانات التي فيها
    غالبًا أعمدة رقمية/تاريخ (رقم تسلسلي، إحداثيات، تواريخ...) تُخفّض نسبة
    "النصية" فيها حتى لو كانت كل خلاياها معبّأة."""
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return False
    import datetime as _dt
    if isinstance(v, (_dt.date, _dt.datetime)):
        return False
    s = str(v).strip()
    if not s:
        return False
    try:
        float(s)
        return False  # نص لكنه رقم صرف مكتوب كنص (مثل رقم إشعار)
    except ValueError:
        return True


def _detect_header_row(ws, max_scan: int = 15) -> int:
    """يخمّن رقم صف رأس الأعمدة الحقيقي بدل افتراض أنه دائمًا أول صف. بعض
    الملفات فيها صف عنوان (مثل "تقرير ملاحظات المغذي" بخلية واحدة مدمجة) قبل
    صف الرؤوس الفعلي. المعيار: صف الرؤوس الحقيقي يكون شبه كامل بتسميات
    نصية عبر كل الأعمدة، بعكس صف عنوان (خلية أو خليتين فقط) أو صف بيانات
    (فيه عادة أعمدة رقمية/تاريخ تُخفّض عدد الخلايا "النصية" فيه) — لذلك نعتمد
    عدد الخلايا النصية الوصفية كمقياس أساسي بدل مجرد عدّ الخلايا غير الفارغة."""
    best_idx, best_score = 0, (-1, -1)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_scan:
            break
        text_like = sum(1 for v in row if _is_text_like(v))
        non_empty = sum(1 for v in row if v is not None and str(v).strip() != "")
        score = (text_like, non_empty)
        if score > best_score:
            best_score, best_idx = score, i
    return best_idx


def list_sheets_with_headers(path: str) -> List[SheetInfo]:
    """يفتح ملف الإكسل (read-only لتفادي استهلاك الذاكرة) ويعيد لكل ورقة
    اسمها وصف رأس الأعمدة (بعد تخمين رقم صف الرؤوس الحقيقي) وعدد الصفوف."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    infos = []
    for name in wb.sheetnames:
        ws = wb[name]
        header_row_idx = _detect_header_row(ws)
        headers: List[str] = []
        n_rows = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == header_row_idx:
                headers = ["" if v is None else str(v) for v in row]
            n_rows += 1
        infos.append(SheetInfo(name=name, header_row_idx=header_row_idx,
                                headers=headers, n_rows=max(n_rows - header_row_idx - 1, 0)))
    wb.close()
    return infos


def suggest_column_mapping(headers: List[str]) -> Dict[str, Optional[int]]:
    """لكل حقل مطلوب، يرجع رقم أفضل عمود مطابق (index بدءًا من صفر) أو None."""
    mapping: Dict[str, Optional[int]] = {}
    used = set()
    for field_key, aliases in FIELD_ALIASES.items():
        best_idx, best_score = None, 0.0
        for idx, h in enumerate(headers):
            if idx in used or not h:
                continue
            for alias in aliases:
                score = similarity(h, alias)
                if score > best_score:
                    best_score, best_idx = score, idx
        if best_idx is not None and best_score >= 0.55:
            mapping[field_key] = best_idx
            used.add(best_idx)
        else:
            mapping[field_key] = None
    return mapping


def read_rows(path: str, sheet_name: str) -> List[List[Any]]:
    """يرجع كل صفوف البيانات (بعد تخطي صف عنوان محتمل وصف الرؤوس الحقيقي) كقوائم خام."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header_row_idx = _detect_header_row(ws)
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= header_row_idx:
            continue
        rows.append(list(row))
    wb.close()
    return rows


def build_records(rows: List[List[Any]], mapping: Dict[str, Optional[int]]) -> List[Dict[str, Any]]:
    """يحوّل الصفوف الخام إلى قواميس بأسماء الحقول المنطقية حسب الربط،
    ويتجاهل الصفوف الفارغة تمامًا."""
    records = []
    for row in rows:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec: Dict[str, Any] = {}
        for field_key, idx in mapping.items():
            if idx is not None and idx < len(row):
                rec[field_key] = row[idx]
            else:
                rec[field_key] = None
        records.append(rec)
    return records


# --------------------------------------------------------------------------
# 2) تحليل قالب البوربوينت: اكتشاف الأشكال الثابتة/المتغيرة
# --------------------------------------------------------------------------

def _emu(v):
    return int(v) if v is not None else 0


def _shape_geo_key(shape, grid=50000):
    """مفتاح هندسي تقريبي لموضع الشكل (لتجميع نفس الدور عبر السلايدات
    حتى لو اختلف الـ id أو الاسم قليلًا)."""
    l = _emu(shape.left) // grid
    t = _emu(shape.top) // grid
    w = _emu(shape.width) // grid
    h = _emu(shape.height) // grid
    return (l, t, w, h)


def _shape_text(shape) -> str:
    if shape.has_text_frame:
        return shape.text_frame.text.strip()
    return ""


@dataclass
class ShapeRole:
    geo_key: Tuple[int, int, int, int]
    shape_type: str            # "text" | "picture"
    samples: List[str] = field(default_factory=list)
    is_variable: bool = False
    left: int = 0
    top: int = 0
    width: int = 0
    height: int = 0
    # أقرب نص "تسمية" ثابت فوق هذا الشكل (إن وُجد) يساعد على التخمين
    nearby_label: str = ""


def analyze_template_slides(prs: Presentation, slide_indices: List[int]) -> List[ShapeRole]:
    """يقارن مجموعة سلايدات عيّنة (نفس نوع سلايد "الملاحظة" المتكرر) ليكتشف
    أي المواضع الهندسية تحمل نصًا/صورة متغيرة من سلايد لآخر، وأيها ثابت."""
    roles: Dict[Tuple[str, Tuple[int, int, int, int]], ShapeRole] = {}

    # ملاحظة: نحلّل الأشكال النصية فقط هنا (ثابت/متغيّر عبر عينات السلايدات).
    # أما الصور فتُعالج بشكل مستقل (get_master_slide_pictures) لأن موضعها
    # يتغيّر قليلًا تلقائيًا عند وجود صورة واحدة فقط بدل اثنتين، مما يُفسد
    # المقارنة الهندسية عبر سلايدات مختلفة.
    all_shapes_by_slide = []
    for si in slide_indices:
        slide = prs.slides[si]
        shapes_here = []
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            stype = "text"
            text = shp.text_frame.text.strip()
            key = _shape_geo_key(shp)
            shapes_here.append((stype, key, text, shp))
        all_shapes_by_slide.append(shapes_here)

    for shapes_here in all_shapes_by_slide:
        for stype, key, text, shp in shapes_here:
            rk = (stype, key)
            if rk not in roles:
                roles[rk] = ShapeRole(
                    geo_key=key, shape_type=stype,
                    left=_emu(shp.left), top=_emu(shp.top),
                    width=_emu(shp.width), height=_emu(shp.height),
                )
            roles[rk].samples.append(text)

    # نص متغيّر = عدد القيم الفريدة (غير الفارغة) أكبر من 1
    for rk, role in roles.items():
        non_empty = [s for s in role.samples if s]
        uniq = set(non_empty)
        role.is_variable = len(uniq) > 1

    result = list(roles.values())
    constant_labels = [r for r in result if not r.is_variable
                        and len(set(s for s in r.samples if s)) == 1 and r.samples[0]]
    _attach_nearby_labels([r for r in result if r.is_variable], constant_labels)
    return result


MAX_LABEL_GAP = 120000        # أقصى مسافة رأسية بين أسفل التسمية وأعلى القيمة (نحو 0.13 بوصة)
MIN_LABEL_OVERLAP_FRAC = 0.3  # أقل نسبة تداخل أفقي مطلوبة بين التسمية والقيمة


def _attach_nearby_labels(target_roles: List[ShapeRole], constant_labels: List[ShapeRole]):
    """يبحث لكل شكل هدف عن أقرب "تسمية" ثابتة تقع فوقه مباشرة (بنفس المحاذاة
    الأفقية تقريبًا) ليستخدمها guess_field_for_role في التخمين."""
    for role in target_roles:
        best_label, best_dist = "", None
        for lbl in constant_labels:
            if lbl.top >= role.top:
                continue
            overlap = min(lbl.left + lbl.width, role.left + role.width) - max(lbl.left, role.left)
            # نقيس نسبة التداخل بالنسبة لعرض "التسمية" نفسها (وليس أصغر عرض بين
            # الاثنين): شكل عريض جدًا يمتد على كامل عرض السلايد (مثل شريط عنوان)
            # قد يتقاطع بالكامل مع شكل هدف ضيّق فيظهر تداخل 100% حسب المقياس
            # القديم (min(lbl.width, role.width)) رغم أنه لا يمثّل تسمية مخصّصة
            # لهذا الشكل تحديدًا. القياس الجديد يرفض هذه الحالة لأن التداخل لا
            # يمثل إلا جزءًا صغيرًا من عرض التسمية العريضة نفسها.
            lbl_width = lbl.width or 1
            if overlap <= 0 or (overlap / lbl_width) < MIN_LABEL_OVERLAP_FRAC:
                continue
            dist = role.top - (lbl.top + lbl.height)
            if -30000 <= dist <= MAX_LABEL_GAP and (best_dist is None or dist < best_dist):
                best_dist = dist
                best_label = lbl.samples[0]
        role.nearby_label = best_label


def get_master_slide_pictures(prs: Presentation, master_index: int,
                               constant_labels: List[ShapeRole]) -> List[ShapeRole]:
    """يرجع أشكال الصور (PICTURE) الموجودة في السلايد النموذجي المُختار فقط
    (وليس عبر عدة سلايدات) لأن هذا هو السلايد الذي سيُستنسخ فعليًا، فموضع
    صوره بالضبط هو ما سيُستخدم لكل السلايدات الجديدة."""
    slide = prs.slides[master_index]
    roles = []
    for shp in slide.shapes:
        if shp.shape_type != 13:
            continue
        role = ShapeRole(
            geo_key=_shape_geo_key(shp), shape_type="picture", samples=[""],
            is_variable=True, left=_emu(shp.left), top=_emu(shp.top),
            width=_emu(shp.width), height=_emu(shp.height),
        )
        roles.append(role)
    _attach_nearby_labels(roles, constant_labels)
    return roles


def pick_best_master_slide(prs: Presentation, slide_indices: List[int]) -> int:
    """يقترح أفضل سلايد ليكون "النموذج" الذي يُستنسخ: الأفضلية للسلايد الذي
    يحتوي أكبر عدد من أشكال الصور (لضمان وجود كل خانات الصور - قبل/بعد - في
    القالب حتى لو كانت بعض السلايدات الأخرى تحتوي صورة واحدة فقط)."""
    best_idx, best_count = slide_indices[0], -1
    for idx in slide_indices:
        n_pics = sum(1 for shp in prs.slides[idx].shapes if shp.shape_type == 13)
        if n_pics > best_count:
            best_count, best_idx = n_pics, idx
    return best_idx


_COORD_RE = re.compile(r"^\s*-?\d{1,3}\.\d{3,},\s*-?\d{1,3}\.\d{3,}\s*$")
_TITLE_NUM_RE = re.compile(r"(\D*)(\d+)(\D*)$")


def guess_field_for_role(role: ShapeRole) -> Tuple[str, str]:
    """يرجع (نوع_الاقتراح, تفاصيل) حيث نوع الاقتراح واحد من:
    'coordinates' | 'seq_title' | 'field:<key>' | 'picture_before' | 'picture_after' | 'picture' | 'unknown'
    """
    sample = next((s for s in role.samples if s), "")
    if role.shape_type == "picture":
        # نستخدم فحص كلمة مفصلية دقيق ("قبل"/"بعد") بدل نسبة تشابه عامة، لأن
        # العبارتين "صورة قبل" و"صورة بعد" متشابهتان جدًا شكليًا (فرق كلمة واحدة
        # فقط) وأي مقياس تشابه ضبابي قد يخلط بينهما.
        lbl_norm = normalize_ar(role.nearby_label)
        has_before = bool(re.search(r"\bقبل\b", lbl_norm))
        has_after = bool(re.search(r"\bبعد\b", lbl_norm))
        if has_before and not has_after:
            return "picture_before", role.nearby_label
        if has_after and not has_before:
            return "picture_after", role.nearby_label
        return "picture", ""
    if _COORD_RE.match(sample):
        return "coordinates", sample
    if all(_TITLE_NUM_RE.match(s or "") for s in role.samples if s):
        # نفس البادئة/اللاحقة مع رقم متغيّر => على الأغلب رقم تسلسلي بعنوان
        prefixes = set()
        suffixes = set()
        for s in role.samples:
            if not s:
                continue
            m = _TITLE_NUM_RE.match(s)
            if m:
                prefixes.add(m.group(1))
                suffixes.add(m.group(3))
        if len(prefixes) == 1 and len(suffixes) == 1:
            prefix = next(iter(prefixes))
            suffix = next(iter(suffixes))
            return "seq_title", f"{prefix}{{seq}}{suffix}"
    # مطابقة عبر التسمية القريبة أو محتوى العينة نفسه مقابل قاموس المرادفات
    label = role.nearby_label
    best_field, best_score = None, 0.0
    for field_key, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            score = similarity(label, alias) if label else 0.0
            if score > best_score:
                best_score, best_field = score, field_key
    if best_field and best_score >= 0.6:
        return f"field:{best_field}", label
    # احتياطي أخير: صندوق نص حر طويل بدون تسمية وبدون نمط معروف => على الأغلب
    # هو نص الملاحظة الرئيسي
    lengths = [len(s) for s in role.samples if s]
    if lengths and (sum(lengths) / len(lengths)) >= 12:
        return "field:note_text", sample
    return "unknown", sample


# --------------------------------------------------------------------------
# 3) استنساخ السلايد وتعبئته
# --------------------------------------------------------------------------

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _copy_slide_relationships(src_slide, dst_slide, new_spTree_element):
    """يجد كل مرجع علاقة (rId) في الشجرة المنسوخة — صور (r:embed/r:link) وأي
    نوع آخر بما فيها الروابط التشعبية hyperlinks (r:id على <a:hlinkClick>)
    — ويعيد ربطه بعلاقة صحيحة داخل السلايد الجديد. تعميم هذا الفحص على كل
    خاصية بمساحة أسماء r: (لا فقط embed/link) ضروري: لو كان السلايد
    "النموذج" الذي نستنسخه يحتوي بالفعل رابطًا تشعبيًا (مثلاً لو استُخدم
    بالخطأ تقرير سابق كنموذج)، فبدون هذا التعميم يبقى الرابط في السلايد
    الجديد يشير إلى rId غير موجود في علاقات السلايد الجديد (مرجع معلّق/تالف)."""
    for el in new_spTree_element.iter():
        for attr_qname in list(el.attrib.keys()):
            if not (attr_qname.startswith("{") and attr_qname[1:].split("}")[0] == _R_NS):
                continue
            rid = el.get(attr_qname)
            if not rid:
                continue
            try:
                rel = src_slide.part.rels[rid]
            except KeyError:
                continue
            if rel.is_external:
                new_rid = dst_slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                try:
                    related_part = src_slide.part.related_part(rid)
                except KeyError:
                    continue
                new_rid = dst_slide.part.relate_to(related_part, rel.reltype)
            el.set(attr_qname, new_rid)


def duplicate_slide(prs: Presentation, index: int):
    """ينسخ سلايد كامل (بكل أشكاله وصوره) ويضيفه في نهاية العرض، ثم يرجع
    الكائن الجديد. لا يغيّر ترتيبه هنا (استخدم move_slide بعد ذلك عند الحاجة)."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    # إزالة أي أشكال أضافها التخطيط (placeholders) تلقائيًا في السلايد الجديد
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)

    # نسخ كل عنصر شكل من الأصل بالترتيب
    for shp in source.shapes:
        new_el = copy.deepcopy(shp._element)
        dest.shapes._spTree.append(new_el)

    _copy_slide_relationships(source, dest, dest.shapes._spTree)
    return dest


def move_slide(prs: Presentation, old_index: int, new_index: int):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])


def delete_slide(prs: Presentation, index: int):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].get(qn("r:id"))
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


def _find_shapes_by_geo(slide, geo_key, grid=50000):
    out = []
    for shp in slide.shapes:
        if _shape_geo_key(shp, grid) == geo_key:
            out.append(shp)
    return out


def set_shape_text_preserve_style(shape, new_text: str):
    """يستبدل نص الشكل مع محاولة الحفاظ على تنسيق أول Run (الخط/الحجم/اللون).

    كما نُفعّل التفاف النص (word_wrap) وتصغير الخط تلقائيًا عند الحاجة
    (auto_size = TEXT_TO_FIT_SHAPE) كشبكة أمان: لا يغيّر هذا شكل النصوص التي
    تناسب الصندوق أصلًا، لكنه يمنع خروج أي نص أطول من المتوقع (شهر فيه
    ملاحظة بوصف أطول من العادة مثلًا) عن حدود صندوقه المخصص بدل أن يفيض
    فوق العناصر المجاورة."""
    tf = shape.text_frame
    try:
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass
    if not tf.paragraphs:
        tf.text = new_text
        return
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = new_text
        # حذف أي runs زائدة في نفس الفقرة
        for r in p0.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.text = new_text
    # حذف أي فقرات إضافية
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def download_image(url: str, timeout=20, retries=2) -> Optional[bytes]:
    if not url or not str(url).startswith("http"):
        return None
    for attempt in range(retries + 1):
        try:
            resp = requests.get(url, timeout=timeout)
            if resp.status_code == 200 and resp.content:
                return resp.content
        except requests.RequestException:
            pass
        time.sleep(0.5)
    return None


def prefetch_images(urls: List[str], max_workers: int = 16, progress_cb=None) -> Dict[str, Optional[bytes]]:
    """ينزّل عدة صور بالتوازي (الشبكة هي القيد وليس المعالج) لتسريع العملية
    بشكل كبير مقارنة بالتنزيل التسلسلي أثناء تعبئة كل سلايد."""
    from concurrent.futures import ThreadPoolExecutor, as_completed
    unique = sorted({u for u in urls if u and str(u).startswith("http")})
    result: Dict[str, Optional[bytes]] = {}
    if not unique:
        return result
    done = 0
    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        futures = {ex.submit(download_image, u): u for u in unique}
        for fut in as_completed(futures):
            u = futures[fut]
            try:
                result[u] = fut.result()
            except Exception:
                result[u] = None
            done += 1
            if progress_cb:
                progress_cb(done, len(unique))
    return result


def insert_picture_in_box(slide, box_left: int, box_top: int, box_w: int, box_h: int, image_bytes: bytes):
    """يضيف صورة *داخل* مساحة محددة بدون أي قص أو تمديد أو تدوير: تصحيح
    اتجاه EXIF أولًا، ثم حساب أبعاد الصورة الحقيقية وتصغيرها/تكبيرها بنفس
    النسبة (Aspect Fit) لتناسب الصندوق، وتوسيطها بداخله."""
    new_left, new_top, new_w, new_h = box_left, box_top, box_w, box_h
    try:
        from PIL import Image, ImageOps
        im = Image.open(io.BytesIO(image_bytes))
        # تصحيح الاتجاه حسب بيانات EXIF (بعض كاميرات الجوال تحفظ الصورة بشكل
        # "مسطّح" مع علم دوران في الميتاداتا)؛ بدون هذا التصحيح تظهر الصورة
        # مستلقية على جنبها. هذا ليس "تدويرًا" نضيفه نحن، بل إعادتها لوضعها
        # الصحيح الذي صُوّرت به أصلًا.
        fixed = ImageOps.exif_transpose(im)
        if fixed is not im:
            im = fixed
            buf = io.BytesIO()
            save_format = "PNG" if im.mode in ("RGBA", "LA", "P") else "JPEG"
            im.convert("RGBA" if save_format == "PNG" else "RGB").save(buf, format=save_format, quality=92)
            image_bytes = buf.getvalue()
        iw, ih = im.size
        if iw and ih and box_w and box_h:
            box_ratio = box_w / box_h
            img_ratio = iw / ih
            if img_ratio > box_ratio:
                new_w = box_w
                new_h = int(round(box_w / img_ratio))
            else:
                new_h = box_h
                new_w = int(round(box_h * img_ratio))
            new_left = box_left + (box_w - new_w) // 2
            new_top = box_top + (box_h - new_h) // 2
    except Exception:
        # في حال تعذّر قراءة أبعاد الصورة نضعها بمقاس الصندوق الأصلي كاحتياط
        pass

    return slide.shapes.add_picture(io.BytesIO(image_bytes), new_left, new_top, new_w, new_h)


def replace_picture_shape(slide, old_shape, image_bytes: bytes):
    """يحذف شكل الصورة القديم ويضيف صورة جديدة داخل نفس مساحته بالضبط
    (بدون قص أو تمديد أو تدوير) عبر insert_picture_in_box."""
    box_left, box_top = old_shape.left, old_shape.top
    box_w, box_h = old_shape.width, old_shape.height
    old_shape._element.getparent().remove(old_shape._element)
    return insert_picture_in_box(slide, box_left, box_top, box_w, box_h, image_bytes)


def load_image_bytes(source: str) -> Optional[bytes]:
    """يقرأ صورة من رابط إنترنت أو من مسار ملف محلي على الجهاز، بحسب الصيغة."""
    if not source:
        return None
    source = source.strip().strip('"').strip("'")
    if source.startswith("http://") or source.startswith("https://"):
        return download_image(source)
    if os.path.exists(source):
        try:
            with open(source, "rb") as f:
                return f.read()
        except OSError:
            return None
    return None


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def set_hyperlink(shape, url: str):
    """يضيف رابطًا (Hyperlink) لكل الأسطر النصية في الشكل، مع إبقاء النص كما هو.
    يبقى الرابط فعّالًا عند تصدير الملف إلى PDF."""
    if not url:
        return
    tf = shape.text_frame
    for p in tf.paragraphs:
        for r in p.runs:
            try:
                r.hyperlink.address = url
            except Exception:
                pass


def maps_url(lat, lon) -> Optional[str]:
    try:
        lat_f = float(str(lat).strip())
        lon_f = float(str(lon).strip())
    except (TypeError, ValueError):
        return None
    return f"https://www.google.com/maps?q={lat_f},{lon_f}"


def detect_coord_order(sample_text: str) -> str:
    """يخمّن ترتيب الإحداثيات في نص مثل '28.04, 42.03' بالاعتماد على نطاق
    خطوط الطول/العرض التقريبي للسعودية (تجنبًا لطلب تأكيد غير ضروري غالبًا).
    يرجع 'lat_lon' أو 'lon_lat'."""
    m = re.match(r"\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)\s*", sample_text or "")
    if not m:
        return "lat_lon"
    a, b = float(m.group(1)), float(m.group(2))
    # نطاق تقريبي للسعودية: خط العرض (lat) ~16-33, خط الطول (lon) ~34-56
    if 10 <= a <= 34 and 30 <= b <= 60:
        return "lat_lon"
    if 30 <= a <= 60 and 10 <= b <= 34:
        return "lon_lat"
    return "lat_lon"


def literal_replace_everywhere(prs: Presentation, old_text: str, new_text: str, slide_indices: Optional[List[int]] = None):
    """يبحث عن نص حرفي (old_text) في كل الشرائح المحددة (أو كل العرض) ويستبدله،
    مع الحفاظ على تنسيق أول Run في كل فقرة تحتوي على النص."""
    if not old_text or old_text == new_text:
        return 0
    count = 0
    indices = slide_indices if slide_indices is not None else range(len(prs.slides._sldIdLst))
    for i in indices:
        slide = prs.slides[i]
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            for p in shp.text_frame.paragraphs:
                full = "".join(r.text for r in p.runs)
                if old_text in full:
                    new_full = full.replace(old_text, new_text)
                    if p.runs:
                        p.runs[0].text = new_full
                        for r in p.runs[1:]:
                            r._r.getparent().remove(r._r)
                        count += 1
    return count


def list_slide_text_shapes(slide) -> List[Dict[str, Any]]:
    """يرجع كل الأشكال النصية غير الفارغة في سلايد واحد (للغلاف/الملخص) مع مواضعها،
    ليستعرضها المستخدم في الويزارد للمراجعة اليدوية."""
    out = []
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip():
            out.append({
                "shape_id": shp.shape_id,
                "geo_key": list(_shape_geo_key(shp)),
                "text": shp.text_frame.text.strip(),
            })
    return out


# --------------------------------------------------------------------------
# 3.5) تحديث سلايد الملخص الإحصائي (Best effort)
# --------------------------------------------------------------------------

_PURE_INT_RE = re.compile(r"^\d+$")


def find_summary_slide(prs: Presentation, keyword: str = "ملخص") -> Optional[int]:
    kw = normalize_ar(keyword)
    for i, slide in enumerate(prs.slides):
        for shp in slide.shapes:
            if shp.has_text_frame and kw in normalize_ar(shp.text_frame.text):
                return i
    return None


def count_pictures_near_box(prs: Presentation, slide_indices: List[int],
                             box: Tuple[int, int, int, int], tolerance: int = 700000) -> int:
    """يحسب عدد السلايدات (ضمن نطاق) التي يوجد فيها بالفعل صورة قريبة من
    صندوق مُعطى (يُستخدم لمعرفة كم ملاحظة أصبح لديها صورة "بعد" الآن)."""
    bl, bt, bw, bh = box
    count = 0
    for idx in slide_indices:
        for shp in prs.slides[idx].shapes:
            if shp.shape_type == 13 and abs(shp.left - bl) <= tolerance and abs(shp.top - bt) <= tolerance:
                count += 1
                break
    return count


def update_summary_totals(prs: Presentation, slide_index: int, total: int,
                           total_caption_kw: str = "اجمالي الملاحظات") -> bool:
    """يبحث عن الرقم المرتبط بتسمية 'إجمالي الملاحظات' (أو ما شابه) ويحدّثه.
    يرجع True إذا نجح الإيجاد والتحديث."""
    slide = prs.slides[slide_index]
    text_shapes = [s for s in slide.shapes if s.has_text_frame]
    kw = normalize_ar(total_caption_kw)
    caption_shape = None
    for s in text_shapes:
        if kw in normalize_ar(s.text_frame.text):
            caption_shape = s
            break
    if caption_shape is None:
        return False
    # أقرب شكل رقمي بحت فوق التسمية مباشرة *ونفس عمودها الأفقي*. لا يكفي
    # الاعتماد على أقرب مسافة رأسية فقط: بطاقات KPI متعددة (مثل "إجمالي
    # الملاحظات" و"الملاحظات المعالجة") غالبًا تتشارك نفس الصف (نفس top)،
    # فتتساوى المسافة الرأسية بينها، وبدون شرط التداخل الأفقي يختار الكود
    # أول صندوق بالصف بدل الصندوق الصحيح فوق هذه التسمية بالضبط — مما قد
    # يحدّث رقمًا خاطئًا (كأن يستبدل "إجمالي الملاحظات" بعدد المعالَجة).
    best, best_dist = None, None
    for s in text_shapes:
        t = s.text_frame.text.strip()
        if not _PURE_INT_RE.match(t):
            continue
        if s.top is None or caption_shape.top is None:
            continue
        if s.top >= caption_shape.top:
            continue
        overlap = (min(s.left + s.width, caption_shape.left + caption_shape.width)
                   - max(s.left, caption_shape.left))
        min_width = min(s.width, caption_shape.width) or 1
        if overlap <= 0 or (overlap / min_width) < 0.3:
            continue
        dist = caption_shape.top - (s.top + s.height)
        if dist >= -30000 and (best_dist is None or dist < best_dist):
            best_dist, best = dist, s
    if best is None:
        return False
    set_shape_text_preserve_style(best, str(total))
    return True


# --------------------------------------------------------------------------
# 3.6) تعبئة سلايد ملاحظة واحدة حسب قرارات الربط
# --------------------------------------------------------------------------

def format_field_value(value: Any) -> str:
    if value is None:
        return ""
    try:
        import datetime as _dt
        if isinstance(value, (_dt.date, _dt.datetime)):
            return value.strftime("%Y-%m-%d")
    except Exception:
        pass
    s = str(value).strip()
    if s.upper() in ("#N/A", "#VALUE!", "NONE", "NAN"):
        return ""
    return s


def fill_slide(slide, decisions: List[Dict[str, Any]], record: Dict[str, Any],
               image_cache: Optional[Dict[str, Optional[bytes]]] = None) -> Dict[str, int]:
    """يطبّق قرارات الربط (decisions) على سلايد مُستنسخ واحد، ويملأه ببيانات
    صف واحد من الإكسل (record). يرجع إحصائية بسيطة عن الصور المنزّلة/الناقصة."""
    if image_cache is None:
        image_cache = {}
    stats = {"images_ok": 0, "images_missing": 0, "before_ok": 0, "after_ok": 0}

    for dec in decisions:
        kind = dec.get("kind")
        if kind == "skip":
            continue

        geo_key = tuple(dec["geo_key"])
        matches = _find_shapes_by_geo(slide, geo_key)
        if not matches:
            continue
        shape = matches[0]

        if kind == "field":
            value = format_field_value(record.get(dec["field_key"]))
            set_shape_text_preserve_style(shape, value)

        elif kind == "seq_title":
            raw = record.get(dec["title_source_field"])
            value = format_field_value(raw)
            template = dec.get("title_template", "{value}")
            text = template.replace("{seq}", value).replace("{value}", value)
            set_shape_text_preserve_style(shape, text)

        elif kind == "coordinates":
            lat = record.get(dec.get("lat_field", "latitude"))
            lon = record.get(dec.get("lon_field", "longitude"))
            lat_s, lon_s = format_field_value(lat), format_field_value(lon)
            text = f"{lat_s}, {lon_s}" if lat_s and lon_s else (lat_s or lon_s)
            set_shape_text_preserve_style(shape, text)
            if dec.get("add_maps_link") and lat_s and lon_s:
                url = maps_url(lat_s, lon_s)
                if url:
                    set_hyperlink(shape, url)

        elif kind in ("picture_before", "picture_after"):
            url = record.get(dec.get("photo_field"))
            url = None if url is None else str(url).strip()
            img_bytes = None
            if url and url.startswith("http"):
                if url not in image_cache:
                    image_cache[url] = download_image(url)
                img_bytes = image_cache[url]
            if img_bytes:
                replace_picture_shape(slide, shape, img_bytes)
                stats["images_ok"] += 1
                stats["before_ok" if kind == "picture_before" else "after_ok"] += 1
            else:
                remove_shape(shape)
                stats["images_missing"] += 1

    return stats


def rebuild_slide_order(prs: Presentation, final_order: List[int], drop_indices: List[int]):
    """يعيد بناء ترتيب السلايدات بالكامل دفعة واحدة (أكفأ من التحريك سلايد
    سلايد)، ويحذف علاقات أي سلايد غير مُدرج في الترتيب الجديد."""
    sldIdLst = prs.slides._sldIdLst
    elements = list(sldIdLst)
    for idx in drop_indices:
        rId = elements[idx].get(qn("r:id"))
        try:
            prs.part.drop_rel(rId)
        except KeyError:
            pass
    for el in elements:
        sldIdLst.remove(el)
    for idx in final_order:
        sldIdLst.append(elements[idx])


def generate_report(prs: Presentation, master_slide_index: int, old_range: Tuple[int, int],
                     records: List[Dict[str, Any]], decisions: List[Dict[str, Any]],
                     progress_cb=None, image_cache: Optional[Dict[str, Optional[bytes]]] = None) -> Dict[str, int]:
    """يولّد سلايدات الملاحظات كاملة: نسخ + تعبئة لكل صف، ثم إزالة نطاق
    السلايدات القديم (بيانات الشهر السابق) ووضع السلايدات الجديدة مكانه بالضبط.

    old_range: (start_index, end_index) شامل الطرفين (0-based) للسلايدات
    القديمة المراد استبدالها بالكامل.
    """
    orig_count = len(prs.slides._sldIdLst)
    if image_cache is None:
        image_cache = {}
    total_ok, total_missing, total_before_ok, total_after_ok = 0, 0, 0, 0

    for i, record in enumerate(records):
        dup = duplicate_slide(prs, master_slide_index)
        stats = fill_slide(dup, decisions, record, image_cache=image_cache)
        total_ok += stats["images_ok"]
        total_missing += stats["images_missing"]
        total_before_ok += stats["before_ok"]
        total_after_ok += stats["after_ok"]
        if progress_cb:
            progress_cb(i + 1, len(records))

    start, end = old_range
    final_order = (list(range(0, start)) +
                   list(range(orig_count, orig_count + len(records))) +
                   list(range(end + 1, orig_count)))
    drop_indices = list(range(start, end + 1))
    rebuild_slide_order(prs, final_order, drop_indices)

    return {
        "generated_slides": len(records),
        "images_ok": total_ok,
        "images_missing": total_missing,
        "before_ok": total_before_ok,
        "after_ok": total_after_ok,
    }


# --------------------------------------------------------------------------
# 3.7) أداة "إضافة صور بعد المعالجة" لاحقًا على ملف تم توليده مسبقًا
# --------------------------------------------------------------------------

def find_picture_boxes(prs: Presentation, slide_indices: List[int]) -> Optional[Dict[str, Tuple[int, int, int, int]]]:
    """يبحث عن أول سلايد ضمن النطاق يحتوي صورتين، ويرجع صندوقي "قبل" و"بعد"
    (الأيمن = قبل، الأيسر = بعد، حسب اتجاه الكتابة من اليمين لليسار)."""
    for idx in slide_indices:
        pics = [shp for shp in prs.slides[idx].shapes if shp.shape_type == 13]
        if len(pics) >= 2:
            pics.sort(key=lambda s: s.left)
            after_shape, before_shape = pics[0], pics[-1]
            return {
                "before": (before_shape.left, before_shape.top, before_shape.width, before_shape.height),
                "after": (after_shape.left, after_shape.top, after_shape.width, after_shape.height),
            }
    return None


def slides_missing_photo(prs: Presentation, slide_indices: List[int],
                          box: Tuple[int, int, int, int], tolerance: int = 700000) -> List[int]:
    """يرجع أرقام السلايدات (0-based) التي لا تحتوي صورة قريبة من الصندوق
    المحدد (أي أن هذه الخانة فاضية فيها)."""
    bl, bt, bw, bh = box
    missing = []
    for idx in slide_indices:
        found = False
        for shp in prs.slides[idx].shapes:
            if shp.shape_type != 13:
                continue
            if abs(shp.left - bl) <= tolerance and abs(shp.top - bt) <= tolerance:
                found = True
                break
        if not found:
            missing.append(idx)
    return missing


def slide_title_preview(slide, max_len: int = 70) -> str:
    """يرجع أفضل نص لتمثيل الملاحظة (يتجنّب عنوان "معالجة الملاحظة رقم..."
    ويفضّل أطول نص وصفي آخر) للعرض على المستخدم أثناء إضافة الصور لاحقًا."""
    texts = [shp.text_frame.text.strip() for shp in slide.shapes
             if shp.has_text_frame and shp.text_frame.text.strip()]
    if not texts:
        return ""
    descriptive = [t for t in texts if "معالجة الملاحظة" not in t and not _PURE_INT_RE.match(t)
                   and not _COORD_RE.match(t)]
    pool = descriptive or texts
    longest = max(pool, key=len)
    return longest[:max_len]


# --------------------------------------------------------------------------
# 4) حفظ / تحميل إعدادات الربط
# --------------------------------------------------------------------------

def save_mapping(path: str, config: Dict[str, Any]):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)


def load_mapping(path: str) -> Optional[Dict[str, Any]]:
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

