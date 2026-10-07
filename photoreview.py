"""مراجعة صور «قبل» و«بعد» بالعين، صفًّا بعد صف.

المشكلة: مولّد التقرير يأخذ «صورة قبل 1» و«صورة بعد 1» دائمًا، وهما كثيرًا
ما لا يقارَنان — «قبل» لقطة عامة للعمود من عشرين مترًا و«بعد» لقطة قريبة
لقاعدته. وفي الصف عادةً ثلاث صور «بعد» أو أكثر، فيها غالبًا واحدة من زاوية
تقارب «قبل» وتُظهر المعالجة.

الحل هنا عين المستخدم لا حكم آلي: تُعرض الملاحظات كلها في صفحة واحدة
للتمرير السريع، وقد اختيرت الصور مبدئيًا بأعلى تشابه بصري، فلا يتدخّل إلا
حيث يرى خطأ. والمخرج نسخة من الإكسل أُعيد فيها ترتيب أعمدة الصور، فتصير
«صورة قبل 1» و«صورة بعد 1» هما ما اختاره — ويستفيد منه مولّد التقرير وكل
أداة أخرى بلا تغيير.

التصغير شرط لا تحسين: المستضيف لا يعطي نسخة مصغّرة (جُرِّبت كل الصيغ)، فكل
صورة نحو 1.1 ميجا — ألف ومئتا صورة لستمئة ملاحظة تعني 1.3 جيجا في المتصفح.
بعد التصغير نحو أربعين كيلوبايت، أي خمسين ميجا للملف كله.
"""
from __future__ import annotations

import concurrent.futures as cf
import hashlib
import io
import os
import re
import urllib.request
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from PIL import Image, ImageOps

import engine
import xlsxedit

# مقاس المصغّرة: يكفي للحكم على الزاوية والمعالجة في التمرير، ويصغّر الحجم
# خمسًا وعشرين مرة. والعرض في الصفحة أصغر منه، فتبقى واضحة عند التكبير.
THUMB = int(os.environ.get("REVIEW_THUMB", "520"))
THUMB_Q = int(os.environ.get("REVIEW_THUMB_Q", "72"))
DL_WORKERS = int(os.environ.get("REVIEW_WORKERS", "12"))
DL_TIMEOUT = 30
UA = "Mozilla/5.0 (compatible; feeder-report/1.0)"

BEFORE_PAT = re.compile(r"^صور[ةه]\s*قبل\s*(\d+)$")
AFTER_PAT = re.compile(r"^صور[ةه]\s*بعد\s*(\d+)$")
PLAIN_PAT = re.compile(r"^(?:ال)?صور[ةه]\s*(\d+)$")

ID_ALIASES = ["رقم الملاحظة", "رقم الملاحظه", "الملاحظة رقم"]
TYPE_ALIASES = ["الملاحظة", "الملاحظه", "وصف الملاحظة"]
INSPECT_ALIASES = ["ملاحظات الفحص", "ملاحظات الفحص الإضافية"]
MAINT_ALIASES = ["ملاحظات الصيانة", "ملاحظات الصيانه"]
CONTRACTOR_ALIASES = ["تمت المعالجة بواسطة", "تمت المعالجه بواسطة", "المقاول"]
STATUS_ALIASES = ["الحالة", "الحاله", "حالة الملاحظة", "حالة الملاحظه"]

# ملاحظات المراجعة الجاهزة، بنصّ المستخدم حرفيًا. والقسمة بين المجموعتين هي
# الافتراض الأوّلي لمصير صورة «بعد» في التقرير — والمستخدم يغيّره بضغطة:
# ما يطعن في الصورة نفسها (موقع آخر، زاوية أخرى) تُحذف معه الصورة ويُكتب
# النص مكانها؛ وما يطعن في المعالجة تبقى معه الصورة لأنها هي الإثبات.
PRESETS = [
    "التأكد من الموقع",
    "الصورة غير مطابقة في الزاوية",
    "توضيح ماذا تم في المعالجة",
    "المعالجة غير واضحة",
    "غير مطابقة للمواصفات",
]
HIDE_BY_DEFAULT = {"التأكد من الموقع", "الصورة غير مطابقة في الزاوية"}
NOTE_SEP = "، "

REVIEW_HEADER = "ملاحظة المراجعة"
PHOTO_HEADER = "صورة المعالجة في التقرير"
PHOTO_HIDDEN = "محذوفة — النص مكانها"
PHOTO_KEPT = "تُعرض مع الملاحظة"

# حالات النظام التي تعني أن المعالجة أُعيدت للمقاول
RETURNED = {"RETURNED", "REJECTED", "مسترجع", "مسترجعة", "معاد", "معادة",
            "مرتجع", "مرتجعة", "مرفوض", "مرفوضة"}


# حالات النظام بأسمائها العربية ونوعها (يحدّد لون الشارة في الصفحة). الأسماء
# نفسها المستعملة في أداة حالة الاعتماد حتى لا تختلف التسمية بين الأداتين.
# و«غير قادر على العمل» حالة مستقلة لا تُدمج في «بانتظار المعالجة»: الفريق
# وصل ولم يستطع، وهذا يحتاج قرارًا لا انتظارًا.
STATUS_INFO: Dict[str, Tuple[str, str]] = {
    "APPROVED": ("تم الاقفال", "ok"),
    "ACCEPTED": ("تم الاقفال", "ok"),
    "CLOSED": ("تم الاقفال", "ok"),
    "DONE": ("بانتظار الاقفال", "done"),
    "COMPLETED": ("بانتظار الاقفال", "done"),
    "FINISHED": ("بانتظار الاقفال", "done"),
    "RETURNED": ("مسترجعة — أُعيدت المعالجة للمقاول", "returned"),
    "REJECTED": ("مسترجعة — أُعيدت المعالجة للمقاول", "returned"),
    "REOPENED": ("مسترجعة — أُعيدت المعالجة للمقاول", "returned"),
    "WORK_NOT_POSSIBLE": ("غير قادر على العمل", "blocked"),
    "IN_PROGRESS": ("جاري العمل", "wait"),
    "ASSIGNED": ("مُسندة — بانتظار المعالجة", "wait"),
    "NEW": ("جديدة — بانتظار المعالجة", "wait"),
    "OPEN": ("جديدة — بانتظار المعالجة", "wait"),
}
_AR_STATUS = {
    "تم الاقفال": "APPROVED", "موافق عليها": "APPROVED", "موافق عليه": "APPROVED",
    "بانتظار الاقفال": "DONE", "تمت المعالجة": "DONE", "منجز": "DONE",
    "مسترجع": "RETURNED", "مسترجعة": "RETURNED", "معاد": "RETURNED",
    "معادة": "RETURNED", "مرتجع": "RETURNED", "مرتجعة": "RETURNED",
    "مرفوض": "RETURNED", "مرفوضة": "RETURNED",
    "غير قادر على العمل": "WORK_NOT_POSSIBLE", "تعذر العمل": "WORK_NOT_POSSIBLE",
    "بانتظار المعالجة": "NEW", "جديد": "NEW", "جديدة": "NEW",
}


def status_info(status: str) -> Dict[str, str]:
    """{"code", "label", "kind"} لحالة الملاحظة كما جاءت من النظام."""
    raw = (status or "").strip()
    if not raw or raw.lower() in ("none", "null", "-"):
        return {"code": "", "label": "بلا حالة", "kind": "none"}
    code = raw.upper().replace(" ", "_")
    if code not in STATUS_INFO:
        code = _AR_STATUS.get(raw) or next(
            (v for k, v in _AR_STATUS.items() if _norm(k) == _norm(raw)), code)
    if code in STATUS_INFO:
        label, kind = STATUS_INFO[code]
        return {"code": code, "label": label, "kind": kind}
    return {"code": raw, "label": raw, "kind": "other"}


def status_key(status: Any) -> str:
    """قيمة الحالة كما في الملف — مفتاح مفاتيح الصور. الفارغة مفتاحها ""."""
    t = str(status or "").strip()
    return "" if t.lower() in ("none", "null", "-") else t


# ترتيب الحالات في صفحة المفاتيح: المقفل أولًا ثم ما بعده في مسار الملاحظة
GROUP_ORDER = ["ok", "done", "returned", "blocked", "wait", "other", "none"]
# اسم مختصر للمفتاح بالتسمية الدارجة في الميدان
GROUP_SHORT = {"ASSIGNED": "جاري المعالجة (مُسندة)", "IN_PROGRESS": "جاري العمل"}


def short_label(status: Any) -> str:
    info = status_info(status_key(status))
    return GROUP_SHORT.get(info["code"]) or info["label"].split(" — ")[0]


def status_groups(an: Dict[str, Any]) -> List[Dict[str, Any]]:
    """حالات الملف بأعدادها: [{key, code, label, short, kind, n, images}]."""
    groups: Dict[str, Dict[str, Any]] = {}
    for n in an["notes"]:
        k = status_key(n.get("status", ""))
        g = groups.get(k)
        if g is None:
            info = status_info(k)
            g = groups[k] = dict(info, key=k, n=0, images=0, short=short_label(k))
        g["n"] += 1
        g["images"] += len(n["before"]) + len(n["after"])
    return sorted(groups.values(),
                  key=lambda g: (GROUP_ORDER.index(g["kind"]) if g["kind"] in GROUP_ORDER
                                 else 99, -g["n"]))


def off_rows(an: Dict[str, Any], off: Any) -> set:
    """صفوف الملاحظات التي أطفأ المستخدم صور حالتها."""
    off = set(off or ())
    if not off:
        return set()
    return {n["row"] for n in an["notes"] if status_key(n.get("status", "")) in off}


def is_returned(status: str) -> bool:
    t = (status or "").strip()
    return t.upper() in RETURNED or _norm(t) in {_norm(x) for x in RETURNED}


def split_note(note: str) -> Tuple[List[str], str]:
    """(الملاحظات الجاهزة، النص الحرّ) من نص محفوظ في الإكسل.

    يعيد بناء حالة الصفحة عند رفع ملف رُوجع من قبل، فلا يضيع عمل الجولة
    السابقة ولا يُكتب فوقه."""
    tags: List[str] = []
    rest: List[str] = []
    for part in [x.strip() for x in (note or "").split(NOTE_SEP.strip()) if x.strip()]:
        (tags if part in PRESETS else rest).append(part)
    return tags, NOTE_SEP.join(rest)


def join_note(tags: List[str], text: str) -> str:
    ordered = [p for p in PRESETS if p in (tags or [])]
    parts = ordered + ([text.strip()] if (text or "").strip() else [])
    return NOTE_SEP.join(parts)


def _norm(s: Any) -> str:
    return engine.normalize_ar(str(s or "").strip())


def _find(headers: List[str], aliases: List[str]) -> Optional[int]:
    want = [_norm(a) for a in aliases]
    for i, h in enumerate(headers):
        if _norm(h) in want:
            return i
    for i, h in enumerate(headers):
        nh = _norm(h)
        if any(w and nh.startswith(w) for w in want):
            return i
    return None


def key_of(url: str) -> str:
    return hashlib.md5(url.strip().encode()).hexdigest()


# --------------------------------------------------------------- التحليل

def analyze(path: str) -> Dict[str, Any]:
    """الملاحظات وصورها من الإكسل، دون تنزيل شيء."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        hdr_i = engine._detect_header_row(ws)
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if hdr_i >= len(rows):
        raise ValueError("لم يُعثر على صف عناوين في الملف.")
    headers = [str(c).strip() if c is not None else "" for c in rows[hdr_i]]

    before_cols: List[Tuple[int, int]] = []
    after_cols: List[Tuple[int, int]] = []
    plain_cols: List[Tuple[int, int]] = []
    for i, h in enumerate(headers):
        t = re.sub(r"\s+", " ", str(h or "").strip())
        m = BEFORE_PAT.match(t)
        if m:
            before_cols.append((int(m.group(1)), i))
            continue
        m = AFTER_PAT.match(t)
        if m:
            after_cols.append((int(m.group(1)), i))
            continue
        m = PLAIN_PAT.match(t)
        if m:
            plain_cols.append((int(m.group(1)), i))
    # الملفات القديمة تسمّي أعمدة «قبل» بـ«صورة 1…» بلا كلمة قبل
    if not before_cols and plain_cols:
        before_cols = plain_cols
    before = [c for _, c in sorted(before_cols)]
    after = [c for _, c in sorted(after_cols)]
    if not before:
        raise ValueError("لم يُعثر على أعمدة صور في الملف.")

    id_c = _find(headers, ID_ALIASES)
    ty_c = _find(headers, TYPE_ALIASES)
    in_c = _find(headers, INSPECT_ALIASES)
    mt_c = _find(headers, MAINT_ALIASES)
    ct_c = _find(headers, CONTRACTOR_ALIASES)
    st_c = _find(headers, STATUS_ALIASES)
    # أعمدة مراجعة من جولة سابقة: تُقرأ ويُكتب فيها بدل فتح عمود مكرر
    rv_c = _find(headers, [REVIEW_HEADER])
    ph_c = _find(headers, [PHOTO_HEADER])

    notes: List[Dict[str, Any]] = []
    for i, row in enumerate(rows):
        if i <= hdr_i:
            continue
        if not any(c is not None and str(c).strip() for c in row):
            continue
        def g(c):
            return str(row[c]).strip() if c is not None and c < len(row) and row[c] is not None else ""
        def urls(cols):
            out = []
            for c in cols:
                v = g(c)
                if v.lower().startswith("http") and v not in out:
                    out.append(v)
            return out
        b, a = urls(before), urls(after)
        if not b and not a:
            continue
        status = g(st_c)
        sinfo = status_info(status)
        tags, text = split_note(g(rv_c))
        notes.append({
            "row": i + 1,                       # رقم الصف في إكسل
            "id": g(id_c), "type": g(ty_c),
            "inspect": g(in_c), "maint": g(mt_c), "contractor": g(ct_c),
            "status": status, "returned": sinfo["kind"] == "returned",
            "blocked": sinfo["kind"] == "blocked",
            "before": b, "after": a,
            # قيم الأعمدة في مواضعها الأصلية (بفراغاتها): المخرج يبادل خليتين
            # فقط، ولو أعاد رصّ القائمة المنظّفة لحذف تكرارًا أو أزاح فراغًا
            "before_raw": [g(c) for c in before],
            "after_raw": [g(c) for c in after],
            "tags": tags, "text": text,
            "hide": g(ph_c).startswith(PHOTO_HIDDEN[:6]) if ph_c is not None else None,
        })
    return {
        "sheet": sheet_name,
        "header_row": hdr_i + 1,
        "headers": headers,
        "n_cols": len(headers),
        "review_col": rv_c,
        "photo_col": ph_c,
        "n_returned": sum(1 for n in notes if n["returned"]),
        "n_blocked": sum(1 for n in notes if n["blocked"]),
        "before_cols": before,
        "after_cols": after,
        "notes": notes,
        "total": len(notes),
        "with_after": sum(1 for n in notes if n["after"]),
        "n_images": sum(len(n["before"]) + len(n["after"]) for n in notes),
    }


# ------------------------------------------------------------- التجهيز

def _fetch(url: str) -> Optional[bytes]:
    try:
        req = urllib.request.Request(url, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=DL_TIMEOUT) as r:
            return r.read()
    except Exception:  # noqa: BLE001
        return None


def _thumb(raw: bytes, dst: str) -> bool:
    try:
        im = ImageOps.exif_transpose(Image.open(io.BytesIO(raw)))
        im = im.convert("RGB")
        im.thumbnail((THUMB, THUMB), Image.LANCZOS)
        im.save(dst, "JPEG", quality=THUMB_Q, optimize=True)
        return True
    except Exception:  # noqa: BLE001
        return False


def prepare(an: Dict[str, Any], thumb_dir: str, progress_cb=None,
            rank: bool = True, only: Optional[set] = None) -> Dict[str, Any]:
    """ينزّل صور الملاحظات ويصغّرها، ثم يرشّح أفضل زوج لكل ملاحظة.

    only: صفوف بعينها (الحالات المفعّلة)؛ صور الحالات المطفأة لا تُنزَّل أصلًا."""
    os.makedirs(thumb_dir, exist_ok=True)
    if only is not None:
        an = dict(an, notes=[n for n in an["notes"] if n["row"] in only])
    urls: List[str] = []
    seen = set()
    for n in an["notes"]:
        for u in n["before"] + n["after"]:
            if u not in seen:
                seen.add(u)
                urls.append(u)

    done = [0]
    total = len(urls)

    def one(u: str) -> Tuple[str, bool]:
        dst = os.path.join(thumb_dir, key_of(u) + ".jpg")
        ok = os.path.exists(dst)
        if not ok:
            raw = _fetch(u)
            ok = bool(raw) and _thumb(raw, dst)
        done[0] += 1
        if progress_cb and done[0] % 5 == 0:
            progress_cb(done[0], total, "download")
        return u, ok

    ok_map: Dict[str, bool] = {}
    with cf.ThreadPoolExecutor(DL_WORKERS) as ex:
        for u, ok in ex.map(one, urls):
            ok_map[u] = ok
    if progress_cb:
        progress_cb(total, total, "download")

    choices: Dict[int, Dict[str, Any]] = {}
    if rank:
        choices = _rank_all(an, thumb_dir, ok_map, progress_cb)
    else:
        for n in an["notes"]:
            b = next((u for u in n["before"] if ok_map.get(u)), "")
            a = next((u for u in n["after"] if ok_map.get(u)), "")
            choices[n["row"]] = {"before": b, "after": a, "score": 0}

    return {
        "images": total,
        "fetched": sum(1 for v in ok_map.values() if v),
        "failed": [u for u, v in ok_map.items() if not v],
        "choices": choices,
        "bytes": sum(os.path.getsize(os.path.join(thumb_dir, f))
                     for f in os.listdir(thumb_dir)),
    }


def _rank_all(an: Dict[str, Any], thumb_dir: str, ok_map: Dict[str, bool],
              progress_cb=None) -> Dict[int, Dict[str, Any]]:
    """يرشّح لكل ملاحظة الزوج الأقرب زاويةً بمطابقة المعالم البصرية.

    المطابقة المطلقة («هل الصورتان لموقع واحد») لا تعمل على هذه الصور —
    قِست فكان وسيط الأزواج الحقيقية 6 نقاط والمخلوطة 5. لكن الترشيح
    *داخل الصف* مسألة أسهل: ثلاثة مرشّحين نرتّبهم، والأعلى غالبًا هو
    الذي صُوّر من زاوية «قبل». والمستخدم يصحّح بضغطة."""
    try:
        import cv2
        import numpy as np
    except Exception:  # noqa: BLE001
        return {n["row"]: {"before": next((u for u in n["before"] if ok_map.get(u)), ""),
                           "after": next((u for u in n["after"] if ok_map.get(u)), ""),
                           "score": 0} for n in an["notes"]}

    sift = cv2.SIFT_create(nfeatures=800)
    bf = cv2.BFMatcher()
    feats: Dict[str, Any] = {}

    def f(u: str):
        if u in feats:
            return feats[u]
        p = os.path.join(thumb_dir, key_of(u) + ".jpg")
        im = cv2.imread(p, cv2.IMREAD_GRAYSCALE) if os.path.exists(p) else None
        feats[u] = sift.detectAndCompute(im, None) if im is not None else None
        return feats[u]

    def score(u1: str, u2: str) -> int:
        a, b = f(u1), f(u2)
        if not a or not b or a[1] is None or b[1] is None:
            return 0
        k1, d1 = a
        k2, d2 = b
        if len(k1) < 8 or len(k2) < 8:
            return 0
        good = [m for m, n in bf.knnMatch(d1, d2, k=2) if m.distance < 0.75 * n.distance]
        if len(good) < 8:
            return len(good)
        src = np.float32([k1[m.queryIdx].pt for m in good]).reshape(-1, 1, 2)
        dst = np.float32([k2[m.trainIdx].pt for m in good]).reshape(-1, 1, 2)
        _M, mask = cv2.findHomography(src, dst, cv2.RANSAC, 6.0)
        return int(mask.sum()) if mask is not None else 0

    out: Dict[int, Dict[str, Any]] = {}
    notes = an["notes"]
    for i, n in enumerate(notes):
        bs = [u for u in n["before"] if ok_map.get(u)]
        as_ = [u for u in n["after"] if ok_map.get(u)]
        if progress_cb and i % 5 == 0:
            progress_cb(i, len(notes), "rank")
        if not bs or not as_:
            out[n["row"]] = {"before": bs[0] if bs else "",
                             "after": as_[0] if as_ else "", "score": 0}
            continue
        best = (-1, bs[0], as_[0])
        for b in bs:
            for a in as_:
                s = score(b, a)
                if s > best[0]:
                    best = (s, b, a)
        out[n["row"]] = {"before": best[1], "after": best[2], "score": best[0]}
        # الصف الواحد قد يحمل عشر صور؛ إفراغ الذاكرة يمنع تضخّمها على خادم صغير
        for u in bs + as_:
            feats.pop(u, None)
    if progress_cb:
        progress_cb(len(notes), len(notes), "rank")
    return out


# --------------------------------------------------------------- المخرج

# عتبة «يصلح للمقارنة»: قِيست على ملف فعلي — دونها لم يصوّر الفريق صورة
# «بعد» من زاوية تقارب «قبل» أصلًا، فلا شيء يُقارَن مهما نظرت
COMPARABLE = int(os.environ.get("REVIEW_COMPARABLE", "12"))


def resolve_hide(ch: Dict[str, Any]) -> bool:
    """مصير صورة «بعد»: اختيار المستخدم إن وُجد، وإلا الافتراض من نوع الملاحظة."""
    if isinstance(ch.get("hide"), bool):
        return ch["hide"]
    return any(t in HIDE_BY_DEFAULT for t in (ch.get("tags") or []))


def write_output(an: Dict[str, Any], choices: Dict[int, Dict[str, Any]],
                 src_path: str, out_path: str, off: Any = None) -> Dict[str, Any]:
    """نسخة من الملف: أعمدة الصور مرتّبة حسب الاختيار، وملاحظات المراجعة.

    الصور لا تُحذف ولا تُضاف: تُبادَل خليتان فقط داخل عمودهما، فتصير
    المختارة في «صورة قبل 1» و«صورة بعد 1». و«احذف الصورة» لا يمسح الرابط
    من الإكسل — الصورة المرفوضة دليلك عند إرجاع العمل — بل يُكتب في عمود
    «صورة المعالجة في التقرير» أنها محذوفة، فيكتب مولّد التقرير النص مكانها.

    off: حالات أطفأ المستخدم صورها — تُمسح روابط صور ملاحظاتها كلها (قبل
    وبعد) مع الرابط التشعبي للخلية، فلا تظهر لها صورة في التقرير."""
    cell_values: Dict[Tuple[int, int], Any] = {}
    moved = 0
    gone = off_rows(an, off)
    by_row = {n["row"]: n for n in an["notes"] if n["row"] not in gone}

    drop_links: set = set()
    move_links: Dict[str, str] = {}
    cleared: Dict[str, int] = {}
    photo_cols = list(an["before_cols"]) + list(an["after_cols"])
    for n in an["notes"]:
        if n["row"] not in gone:
            continue
        lbl = short_label(n.get("status", ""))
        cleared[lbl] = cleared.get(lbl, 0) + 1
        for c in photo_cols:
            cell_values[(n["row"], c)] = None
            drop_links.add("%s%d" % (xlsxedit.col_letter(c), n["row"]))

    for row, ch in choices.items():
        n = by_row.get(row)
        if not n:
            continue
        for kind, cols in (("before", an["before_cols"]), ("after", an["after_cols"])):
            raw = list(n.get(kind + "_raw") or [])
            pick = (ch.get(kind) or "").strip()
            if not raw or not pick or raw[0].strip() == pick:
                continue
            j = next((k for k, v in enumerate(raw) if v.strip() == pick), None)
            if j is None:
                continue
            raw[0], raw[j] = raw[j], raw[0]
            cell_values[(row, cols[0])] = raw[0] or None
            cell_values[(row, cols[j])] = raw[j] or None
            a_ref = "%s%d" % (xlsxedit.col_letter(cols[0]), row)
            b_ref = "%s%d" % (xlsxedit.col_letter(cols[j]), row)
            move_links[a_ref], move_links[b_ref] = b_ref, a_ref
            moved += 1

    # ملاحظات المراجعة: في أعمدة الجولة السابقة إن وُجدت، وإلا عمودان جديدان
    n_cols = an["n_cols"]
    new_columns: List[Tuple[str, Dict[int, Any]]] = []
    rv_vals: Dict[int, Any] = {}
    ph_vals: Dict[int, Any] = {}
    n_notes = n_hidden = n_ret_noted = 0
    for n in an["notes"]:
        if n["row"] in gone:
            continue
        ch = choices.get(n["row"]) or {}
        tags = ch.get("tags") if "tags" in ch else n.get("tags")
        text = ch.get("text") if "text" in ch else n.get("text")
        note = join_note(tags or [], text or "")
        if note:
            n_notes += 1
            n_ret_noted += 1 if n.get("returned") else 0
            hide = resolve_hide({"tags": tags, "hide": ch.get("hide", n.get("hide"))})
            if not n.get("after"):
                photo = "لا توجد صورة «بعد» — النص مكانها"
            elif hide:
                photo = PHOTO_HIDDEN
                n_hidden += 1
            else:
                photo = PHOTO_KEPT
            rv_vals[n["row"]] = note
            ph_vals[n["row"]] = photo
        elif an.get("review_col") is not None:
            # أُزيلت ملاحظة كانت في جولة سابقة: تُمسح لا تُترك قديمة
            rv_vals[n["row"]] = None
            ph_vals[n["row"]] = None

    def place(existing: Optional[int], header: str, vals: Dict[int, Any]) -> int:
        if existing is not None:
            for r, v in vals.items():
                cell_values[(r, existing)] = v
            return existing
        idx = n_cols + len(new_columns)
        new_columns.append((header, {r: v for r, v in vals.items() if v}))
        return idx

    cell_fills: Dict[Tuple[int, int], str] = {}
    if rv_vals:
        rc = place(an.get("review_col"), REVIEW_HEADER, rv_vals)
        pc = place(an.get("photo_col"), PHOTO_HEADER, ph_vals)
        # تُلوَّن خليتا الملاحظة وحدهما لا الصف: لون الصف عند المستخدم بيانات
        for r, v in rv_vals.items():
            if v:
                cell_fills[(r, rc)] = "FFC7CE"
                cell_fills[(r, pc)] = "FFC7CE"

    xlsxedit.write_patched(
        src_path, out_path,
        sheet_name=an["sheet"], header_row=an["header_row"],
        n_cols=n_cols, cell_values=cell_values, new_columns=new_columns,
        cell_fills=cell_fills, drop_links=drop_links, move_links=move_links)

    comparable = sum(1 for r, c in choices.items()
                     if by_row.get(r) and by_row[r]["after"]
                     and c.get("score", 0) >= COMPARABLE)
    with_after = sum(1 for n in by_row.values() if n["after"])
    return {
        "notes": len(by_row),
        "with_after": with_after,
        "no_after": len(by_row) - with_after,
        "reordered": moved,
        "comparable": comparable,
        "not_comparable": with_after - comparable,
        "review_notes": n_notes,
        "hidden": n_hidden,
        "returned": sum(1 for n in by_row.values() if n.get("returned")),
        "photos_cleared": len(gone),
        "photos_cleared_by": sorted(cleared.items(), key=lambda t: -t[1]),
        "returned_noted": n_ret_noted,
        "stem": os.path.splitext(os.path.basename(src_path))[0],
    }
