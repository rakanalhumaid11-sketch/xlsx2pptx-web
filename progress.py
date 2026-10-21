"""المنجز والمتبقي: جدول واحد لكل المغذيات، ووراءه ملف كل مغذٍّ كما هو.

المدخل ملفات المغذيات بعد أداة حالة الاعتماد وبعد تلوين المستخدم:
* الأخضر مقفل،
* الأحمر (والأزرق معه) متبقٍّ يحتاج فصل جهد متوسط MV،
* بلا لون متبقٍّ جهد منخفض يمكن العمل عليه.

المخرج ملف إكسل: الورقة الأولى جدول فيه لكل مغذٍّ الإجمالي والمنجز ونسبة
الإنجاز والمتبقي جهد متوسط والمتبقي جهد منخفض، واسم المغذي فيها رابط إلى
ورقته. وبعدها ورقة لكل مغذٍّ منسوخة كما هي (merge.py) — الملف نفسه محفوظ
داخل الملخص، والوصول إليه بضغطة.

اللون المعتمد هو الظاهر في إكسل، لأنه ما يعرّفه المستخدم ويعدّه بعينه:
أداة الاعتماد تلوّن المقفل أخضر بتنسيق شرطي فوق لون الصف الأصلي، فالصف
الأحمر الذي أُقفل يظهر أخضر ويُعدّ منجزًا. أما الوردي الذي تضعه للمسترجع
فليس من ألوان المستخدم: يُقرأ اللون الذي تحته، وهو تصنيفه هو للملاحظة.
"""
from __future__ import annotations

import datetime
import os
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

import merge
import sorter
import summary
from xlsxedit import SheetBuilder, Styles

SHEET_TITLE = "المنجز والمتبقي"
TOOL_PALETTE = summary.TOOL_PALETTE
ID_ALIASES = ["رقم الملاحظة", "رقم الملاحظه", "الملاحظة رقم"]

# الأعمدة كما طلبها المستخدم، بلون رأس يطابق لون الصفوف في ملفاته
COLS = [
    ("total", "الإجمالي", "1F3864", "FFFFFF"),
    ("done", "المنجز", "C6EFCE", "006100"),
    ("pct", "نسبة الإنجاز", "1F3864", "FFFFFF"),
    ("mv", "المتبقي جهد متوسط (يحتاج فصل MV)", "FF0000", "FFFFFF"),
    ("lv", "المتبقي جهد منخفض (ممكن العمل عليه)", "F2F2F2", "000000"),
]
OTHER_COL = ("other", "ألوان أخرى", "FFF2CC", "7F6000")

# ألوان الرسم: الأخضر والأحمر بمعنى المستخدم نفسه، والرمادي لـ«بلا لون».
# الأخضر والأحمر متباعدان بما يكفي لعمى الألوان الأحمر-الأخضر (قِيس: فرق 8.4)،
# ومع ذلك كل شريحة عليها رقمها وبين الشرائح فاصل أبيض، فلا يعتمد الفهم على اللون
CHART_SERIES = [
    ("done", "المنجز", "1F8A5A", "FFFFFF"),
    # بلا «MV» في مفتاح الرسم: كلمة لاتينية في آخر نص عربي تقفز إلى أوله
    ("mv", "متبقٍّ جهد متوسط — يحتاج فصل", "D9483B", "FFFFFF"),
    ("lv", "متبقٍّ جهد منخفض — ممكن العمل عليه", "A6A6A6", "262626"),
]
CHART_OTHER = ("other", "ألوان أخرى", "D9A520", "262626")


def bucket(shown: str, static: str) -> str:
    """done / mv / lv / other لملاحظة واحدة من لونها الظاهر ولونها المخزَّن."""
    shown = (shown or "").upper()
    static = (static or "").upper()
    if sorter.classify(shown) == "green":
        return "done"
    # لون أداة الاعتماد (وردي المسترجع، أصفر بانتظار الإقفال) يغطي لون المستخدم:
    # التصنيف من اللون الذي تحته
    base = shown if shown and shown not in TOOL_PALETTE else static
    if base in TOOL_PALETTE:
        base = ""
    k = sorter.classify(base)
    if k in ("red", "blue"):
        return "mv"
    if k in ("none", "green"):
        # الأخضر تحت تنسيق آخر (مسترجع مثلًا) ليس مقفلًا: يُعدّ ممكن العمل عليه
        return "lv"
    return "other"


def read_file(path: str, name: str) -> Dict[str, Any]:
    """أعداد ملف واحد: الإجمالي والمنجز والمتبقي بنوعيه."""
    an = sorter.analyze(path)
    # الصفوف التي لها رقم ملاحظة فقط: سطر إجمالي أو ملاحظة حرّة أسفل الجدول لا يُعدّ
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[an["sheet"]]
        idc = None
        ids: Dict[int, str] = {}
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == an["header_row"]:
                header = [str(c or "").strip() for c in row]
                idc = next((header.index(a) for a in ID_ALIASES if a in header), None)
                if idc is None:
                    break
                continue
            if idc is None or i < an["header_row"]:
                continue
            v = row[idc] if idc < len(row) else None
            if v is not None and str(v).strip():
                ids[i] = str(v).strip()
    finally:
        wb.close()

    counts = {"total": 0, "done": 0, "mv": 0, "lv": 0, "other": 0}
    for r in an["records"]:
        if ids and r["row"] not in ids:
            continue
        counts["total"] += 1
        counts[bucket(r.get("shown", ""), r.get("hex", ""))] += 1
    return {"name": name, "path": path, **counts}


def _styles(st: Styles) -> Dict[str, int]:
    b = st.border("BFBFBF")
    f = st.font

    def x(**kw):
        kw.setdefault("rtl", True)
        return st.xf(**kw)
    s = {
        "title": x(font=f(16, True, "FFFFFF", "Arial"), fill=st.fill("1F3864"),
                   halign="center", wrap=True),
        "sub": x(font=f(10, False, "595959", "Arial"), halign="center", wrap=True),
        "code": x(font=f(11, True, "0563C1", "Arial"), border=b, halign="center"),
        "num": x(font=f(11, False, "000000", "Arial"), border=b, halign="center"),
        "pct": x(font=f(11, True, "000000", "Arial"), border=b, numfmt=9, halign="center"),
        "note": x(font=f(10, False, "7F7F7F", "Arial"), halign="right", wrap=True),
    }
    sum_fill = st.fill("D9E1F2")
    s["sum"] = x(font=f(11, True, "000000", "Arial"), fill=sum_fill, border=b, halign="center")
    s["sum_pct"] = x(font=f(11, True, "000000", "Arial"), fill=sum_fill, border=b,
                     numfmt=9, halign="center")
    s["th_code"] = x(font=f(11, True, "FFFFFF", "Arial"), fill=st.fill("1F3864"),
                     border=b, halign="center", wrap=True)
    for key, _label, bg, fg in COLS + [OTHER_COL]:
        s["th_" + key] = x(font=f(11, True, fg, "Arial"), fill=st.fill(bg),
                           border=b, halign="center", wrap=True)
    return s


def build_front(rows: List[Dict[str, Any]], date_text: str):
    """دالة تبني الورقة الأولى بعد أن يعرف الدمج أسماء أوراق المغذيات."""
    show_other = any(r["other"] for r in rows)
    cols = COLS + ([OTHER_COL] if show_other else [])

    def front(st: Styles, sources: List[Any]) -> str:
        s = _styles(st)
        sheet_of = {src.path: src.sheet_name for src in sources}
        sb = SheetBuilder(rtl=True, gridlines=False, selected=True, centered=True,
                          tab_color="1F3864")
        last = 1 + len(cols)                      # B المغذي، ثم الأعمدة
        sb.width(0, 3)
        sb.width(1, 18)
        for j, (key, *_r) in enumerate(cols):
            sb.width(2 + j, 15 if key in ("total", "done", "pct") else 22)
        sb.set(1, 1, "المنجز والمتبقي حسب المغذي", s["title"])
        sb.merge(1, 1, 1, last)
        sb.height(1, 34)
        # التاريخ آخر السطر ومسبوق بكلمة: رقم في أول سطر عربي يقلبه محرّك الاتجاه
        sb.set(2, 1, "بيانات ملفات المغذيات بتاريخ %s" % date_text, s["sub"])
        sb.merge(2, 1, 2, last)
        sb.height(2, 20)

        hr = 4
        sb.set(hr, 1, "المغذي", s["th_code"])
        for j, (key, label, *_c) in enumerate(cols):
            sb.set(hr, 2 + j, label, s["th_" + key])
        sb.height(hr, 36)

        first = hr + 1
        for i, r in enumerate(rows):
            row = first + i
            name = sheet_of.get(r["path"], r["name"])
            sb.set(row, 1, name, s["code"])
            # اسم المغذي رابط إلى ورقته في الملف نفسه
            sb.link(row, 1, "%s!A1" % merge.quote_sheet(name))
            for j, (key, *_c) in enumerate(cols):
                col = 2 + j
                L = merge.col_letter
                if key == "pct":
                    sb.set(row, col, (r["done"] / r["total"]) if r["total"] else 0.0, s["pct"],
                           formula="IFERROR(%s%d/%s%d,0)" % (L(3), row, L(2), row))
                else:
                    sb.set(row, col, r[key], s["num"])
            sb.height(row, 20)
        lastrow = first + len(rows) - 1
        tot = lastrow + 1
        sb.set(tot, 1, "الإجمالي", s["sum"])
        for j, (key, *_c) in enumerate(cols):
            col = 2 + j
            L = merge.col_letter(col)
            if key == "pct":
                t, d = sum(r["total"] for r in rows), sum(r["done"] for r in rows)
                sb.set(tot, col, (d / t) if t else 0.0, s["sum_pct"],
                       formula="IFERROR(%s%d/%s%d,0)" % (merge.col_letter(3), tot,
                                                          merge.col_letter(2), tot))
            else:
                sb.set(tot, col, sum(r[key] for r in rows), s["sum"],
                       formula="SUM(%s%d:%s%d)" % (L, first, L, lastrow))
        sb.height(tot, 22)
        pct_col = merge.col_letter(2 + [k for k, *_r in cols].index("pct"))
        sb.databar("%s%d:%s%d" % (pct_col, first, pct_col, lastrow), "5B9BD5")

        # الرسم تحت الجدول بعرضه، وطوله على عدد المغذيات
        col_of = {k: 2 + j for j, (k, *_r) in enumerate(cols)}
        series = [se for se in CHART_SERIES + ([CHART_OTHER] if show_other else [])]
        names = [sheet_of.get(r["path"], r["name"]) for r in rows]
        chart = bar_chart(SHEET_TITLE, names, "$B$%d:$B$%d" % (first, lastrow), [
            {"name": label, "color": color, "label": ink,
             "ref": "$%s$%d:$%s$%d" % (merge.col_letter(col_of[key]), first,
                                       merge.col_letter(col_of[key]), lastrow),
             "values": [r[key] for r in rows]}
            for key, label, color, ink in series])
        top = tot + 1                               # صف فارغ بين الجدول والرسم (ترقيم من صفر)
        height = max(16, 7 + 3 * len(rows))
        sb.drawing_rid = "rId1"
        return sb.to_xml(), [(chart, (1, top, last + 1, top + height))]

    return front


def _txpr(size: int, bold: bool, color: str) -> str:
    return ('<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="%d" b="%d">'
            '<a:solidFill><a:srgbClr val="%s"/></a:solidFill><a:latin typeface="Arial"/>'
            '<a:cs typeface="Arial"/></a:defRPr></a:pPr><a:endParaRPr lang="ar-SA"/></a:p>'
            "</c:txPr>" % (size, 1 if bold else 0, color))


def bar_chart(sheet: str, categories: List[str], cat_ref: str,
              series: List[Dict[str, Any]]) -> str:
    """أشرطة أفقية متراكبة، شريط لكل مغذٍّ وأوله في الأعلى، وعلى كل شريحة رقمها.

    القيم المحسوبة مخزّنة في الرسم نفسه فيظهر كاملًا في عارض الجوال، وهو
    مربوط بخلايا الجدول فيتحدّث معها في إكسل. محور القيم مخفي: الأرقام على
    الشرائح تغني عنه وتُبقي الرسم نظيفًا."""
    q = merge.quote_sheet(sheet)
    esc = merge.esc

    def str_cache(items: List[str]) -> str:
        return '<c:strCache><c:ptCount val="%d"/>%s</c:strCache>' % (
            len(items), "".join('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, esc(v))
                                for i, v in enumerate(items)))

    def num_cache(vals: List[float]) -> str:
        return ('<c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="%d"/>%s'
                "</c:numCache>" % (len(vals), "".join(
                    '<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, v) for i, v in enumerate(vals))))

    sers = []
    for i, se in enumerate(series):
        sers.append(
            '<c:ser><c:idx val="%d"/><c:order val="%d"/>'
            "<c:tx><c:v>%s</c:v></c:tx>"
            '<c:spPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:ln w="19050"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln></c:spPr>'
            '<c:invertIfNegative val="0"/>'
            # الصفر لا يُكتب: شريحة بلا عرض لا تحتاج رقمًا فوقها
            '<c:dLbls><c:numFmt formatCode="0;;;" sourceLinked="0"/>'
            "<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>%s"
            '<c:dLblPos val="ctr"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
            '<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/>'
            '<c:showBubbleSize val="0"/></c:dLbls>'
            "<c:cat><c:strRef><c:f>%s!%s</c:f>%s</c:strRef></c:cat>"
            "<c:val><c:numRef><c:f>%s!%s</c:f>%s</c:numRef></c:val></c:ser>" % (
                i, i, esc(se["name"]), se["color"], _txpr(1100, True, se["label"]),
                q, cat_ref, str_cache(categories), q, se["ref"], num_cache(se["values"])))

    title = ('<c:title><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
             '<a:defRPr sz="1500" b="1"><a:solidFill><a:srgbClr val="1F3864"/></a:solidFill>'
             '<a:latin typeface="Arial"/><a:cs typeface="Arial"/></a:defRPr></a:pPr>'
             '<a:r><a:rPr lang="ar-SA" sz="1500" b="1"><a:solidFill><a:srgbClr val="1F3864"/>'
             '</a:solidFill><a:latin typeface="Arial"/><a:cs typeface="Arial"/></a:rPr>'
             "<a:t>المنجز والمتبقي حسب المغذي</a:t></a:r></a:p></c:rich></c:tx>"
             '<c:overlay val="0"/></c:title>')
    axis_ln = ('<c:spPr><a:noFill/><a:ln w="9525"><a:solidFill><a:srgbClr val="BFBFBF"/>'
               "</a:solidFill></a:ln></c:spPr>")
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<c:chartSpace xmlns:c="%s" xmlns:a="%s" xmlns:r="%s">'
        '<c:date1904 val="0"/><c:lang val="ar-SA"/><c:roundedCorners val="0"/>'
        "<c:chart>%s"
        '<c:autoTitleDeleted val="0"/><c:plotArea><c:layout/>'
        '<c:barChart><c:barDir val="bar"/><c:grouping val="stacked"/><c:varyColors val="0"/>'
        '%s<c:gapWidth val="55"/><c:overlap val="100"/>'
        '<c:axId val="51236301"/><c:axId val="51236302"/></c:barChart>'
        # أول مغذٍّ في الأعلى كما في الجدول
        '<c:catAx><c:axId val="51236301"/><c:scaling><c:orientation val="maxMin"/></c:scaling>'
        '<c:delete val="0"/><c:axPos val="l"/><c:numFmt formatCode="General" sourceLinked="1"/>'
        '<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
        '%s%s<c:crossAx val="51236302"/><c:crosses val="autoZero"/><c:auto val="1"/>'
        '<c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'
        # الأشرطة تبدأ من اليمين كقراءة العربية، وأسماء المغذيات على اليمين
        '<c:valAx><c:axId val="51236302"/><c:scaling><c:orientation val="maxMin"/>'
        '<c:min val="0"/></c:scaling><c:delete val="1"/><c:axPos val="b"/>'
        '<c:numFmt formatCode="General" sourceLinked="1"/><c:majorTickMark val="none"/>'
        '<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
        '<c:crossAx val="51236301"/><c:crosses val="autoZero"/><c:crossBetween val="between"/>'
        "</c:valAx><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>"
        '<c:legend><c:legendPos val="b"/><c:overlay val="0"/>%s</c:legend>'
        '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
        '<c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
        '<a:ln w="9525"><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill></a:ln></c:spPr>'
        "</c:chartSpace>" % (
            merge.C_NS, merge.A_NS, merge.REL_NS, title, "".join(sers),
            axis_ln, _txpr(1100, True, "1F3864"), _txpr(1000, False, "404040")))


def build(files: List[Tuple[str, str]], out_path: str, today: str = "") -> Dict[str, Any]:
    """files: [(المسار، الاسم)] ← ملف الملخص مع أوراق المغذيات، وأرقامه."""
    rows = [read_file(p, n) for p, n in files]
    date_text = today if _is_date(today) else summary.today_riyadh()
    # ترتيب الجدول نفسه ترتيب أوراق المغذيات في الدمج (بالرمز)
    ordered = _sorted_files(files)
    by_path = {r["path"]: r for r in rows}
    res = merge.merge(files, out_path, include_data=False,
                      front=(SHEET_TITLE, build_front([by_path[p] for p, _n in ordered],
                                                      date_text)))
    sheet_rows = []
    for (p, _n), sh in zip(ordered, res["sheets"]):
        r = dict(by_path[p])
        r["sheet"] = sh["name"]
        sheet_rows.append(r)
    warnings = list(res["warnings"])
    n_other = sum(r["other"] for r in rows)
    if n_other:
        warnings.append("%d ملاحظة بألوان غير الأخضر والأحمر والأزرق — في عمود «ألوان أخرى»."
                        % n_other)
    return {"rows": sheet_rows, "date": date_text, "warnings": warnings,
            "other": bool(n_other)}


# ترتيب المغذيات: merge.merge يرتّب الأوراق بالرمز، والجدول يتبعها
def _sorted_files(files: List[Tuple[str, str]]) -> List[Tuple[str, str]]:
    codes = {}
    for p, n in files:
        try:
            src = merge.Source(p, n)
            codes[p] = src.code or n
            src.close()
        except Exception:  # noqa: BLE001
            codes[p] = n
    return sorted(files, key=lambda f: sorter.natural_key(codes[f[0]]))


def _is_date(t: str) -> bool:
    try:
        datetime.datetime.strptime((t or "")[:10], "%Y-%m-%d")
        return len(t) == 10
    except ValueError:
        return False
